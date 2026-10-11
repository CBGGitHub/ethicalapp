"""Convert between Word .docx files and Qt QTextDocuments (used by the word processor)."""
import io
import re
from dataclasses import dataclass

import docx
from docx.enum.dml import MSO_COLOR_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt, QUrl
from PySide6.QtGui import (QBrush, QColor, QFont, QImage, QTextBlockFormat, QTextCharFormat,
                           QTextCursor, QTextDocument, QTextFormat, QTextImageFormat, QTextLength,
                           QTextListFormat, QTextTableCellFormat, QTextTableFormat)

PX_PER_PT = 96 / 72
EMU_PER_PX = 9525
TITLE_PROP = QTextFormat.UserProperty + 1          # block is a "Title" paragraph

# heading level -> (size pt, colour) - roughly Word's built-in look
HEADING_LOOK = {0: (28, "#000000"), 1: (16, "#2F5496"), 2: (13, "#2F5496"), 3: (12, "#1F3763"),
                4: (11, "#2F5496"), 5: (11, "#2F5496"), 6: (11, "#1F3763")}

HIGHLIGHT = {
    WD_COLOR_INDEX.YELLOW: "#FFFF00", WD_COLOR_INDEX.BRIGHT_GREEN: "#00FF00",
    WD_COLOR_INDEX.TURQUOISE: "#00FFFF", WD_COLOR_INDEX.PINK: "#FF00FF",
    WD_COLOR_INDEX.BLUE: "#0000FF", WD_COLOR_INDEX.RED: "#FF0000",
    WD_COLOR_INDEX.DARK_BLUE: "#000080", WD_COLOR_INDEX.TEAL: "#008080",
    WD_COLOR_INDEX.GREEN: "#008000", WD_COLOR_INDEX.VIOLET: "#800080",
    WD_COLOR_INDEX.DARK_RED: "#800000", WD_COLOR_INDEX.DARK_YELLOW: "#808000",
    WD_COLOR_INDEX.GRAY_50: "#808080", WD_COLOR_INDEX.GRAY_25: "#C0C0C0",
    WD_COLOR_INDEX.BLACK: "#000000", WD_COLOR_INDEX.WHITE: "#FFFFFF",
}
HIGHLIGHT_BY_HEX = {v: k for k, v in HIGHLIGHT.items()}

ALIGN_IN = {WD_ALIGN_PARAGRAPH.CENTER: Qt.AlignHCenter, WD_ALIGN_PARAGRAPH.RIGHT: Qt.AlignRight,
            WD_ALIGN_PARAGRAPH.JUSTIFY: Qt.AlignJustify, WD_ALIGN_PARAGRAPH.DISTRIBUTE: Qt.AlignJustify}

LIST_FMT_IN = {"bullet": QTextListFormat.ListDisc, "decimal": QTextListFormat.ListDecimal,
               "lowerLetter": QTextListFormat.ListLowerAlpha,
               "upperLetter": QTextListFormat.ListUpperAlpha,
               "lowerRoman": QTextListFormat.ListLowerRoman,
               "upperRoman": QTextListFormat.ListUpperRoman}
BULLET_STYLES = (QTextListFormat.ListDisc, QTextListFormat.ListCircle, QTextListFormat.ListSquare)
LIST_FMT_OUT = {QTextListFormat.ListDecimal: "decimal", QTextListFormat.ListLowerAlpha: "lowerLetter",
                QTextListFormat.ListUpperAlpha: "upperLetter",
                QTextListFormat.ListLowerRoman: "lowerRoman",
                QTextListFormat.ListUpperRoman: "upperRoman"}
BULLET_CHARS = {QTextListFormat.ListDisc: "•", QTextListFormat.ListCircle: "o",
                QTextListFormat.ListSquare: "▪"}


@dataclass
class PageSetup:
    width_in: float = 8.5
    height_in: float = 11.0
    margin_left: float = 1.0
    margin_top: float = 1.0
    margin_right: float = 1.0
    margin_bottom: float = 1.0
    header: str = ""
    footer: str = ""
    page_numbers: bool = False

    @property
    def size_px(self):
        return self.width_in * 96, self.height_in * 96


def image_bytes(img, fmt="PNG"):
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.WriteOnly)
    img.save(buf, fmt)
    return bytes(ba)


# ====================================================================== import

class _Importer:
    def __init__(self, d, qdoc):
        self.d, self.qdoc = d, qdoc
        self.img_count = 0
        self.lists = {}           # (numId, ilvl) -> QTextList
        self.minor_font, self.major_font = self._theme_fonts()
        self.default_font, self.default_size = self._defaults()
        self.def_after, self.def_line = self._default_spacing()

    # ---- document defaults
    def _theme_fonts(self):
        minor, major = "Calibri", "Calibri Light"
        try:
            for rel in self.d.part.rels.values():
                if rel.reltype == RT.THEME:
                    from lxml import etree
                    root = etree.fromstring(rel.target_part.blob)
                    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                    mi = root.find(".//a:minorFont/a:latin", ns)
                    ma = root.find(".//a:majorFont/a:latin", ns)
                    minor = mi.get("typeface") if mi is not None else minor
                    major = ma.get("typeface") if ma is not None else major
        except Exception:
            pass
        return minor, major

    def _defaults(self):
        family, size = None, None
        dd = self.d.styles.element.find(qn("w:docDefaults"))
        rpr = dd.find(qn("w:rPrDefault") + "/" + qn("w:rPr")) if dd is not None else None
        if rpr is not None:
            family = self._rfonts_name(rpr.find(qn("w:rFonts")))
            sz = rpr.find(qn("w:sz"))
            if sz is not None:
                size = int(sz.get(qn("w:val"))) / 2
        normal = self._style("Normal")
        if normal is not None:
            family = self._style_font(normal) or family
            size = self._style_attr(normal, "size", font=True)
            size = size.pt if size else None
        return family or self.minor_font, size or (11 if dd is not None else 10)

    def _default_spacing(self):
        after, line = 0, None
        dd = self.d.styles.element.find(qn("w:docDefaults"))
        if dd is not None:
            sp = dd.find(f"{qn('w:pPrDefault')}/{qn('w:pPr')}/{qn('w:spacing')}")
            if sp is not None:
                if sp.get(qn("w:after")):
                    after = int(sp.get(qn("w:after"))) / 20
                if sp.get(qn("w:line")) and sp.get(qn("w:lineRule"), "auto") == "auto":
                    line = int(sp.get(qn("w:line"))) / 240
        return after, line

    def _style(self, name):
        try:
            return self.d.styles[name]
        except KeyError:
            return None

    def _rfonts_name(self, rfonts):
        if rfonts is None:
            return None
        name = rfonts.get(qn("w:ascii")) or rfonts.get(qn("w:hAnsi"))
        if name:
            return name
        theme = rfonts.get(qn("w:asciiTheme")) or rfonts.get(qn("w:hAnsiTheme"))
        if theme:
            return self.major_font if theme.startswith("major") else self.minor_font
        return None

    def _style_font(self, style):
        s = style
        while s is not None:
            rpr = s.element.rPr
            if rpr is not None:
                name = self._rfonts_name(rpr.find(qn("w:rFonts")))
                if name:
                    return name
            s = s.base_style
        return None

    @staticmethod
    def _style_attr(style, attr, font=False):
        s = style
        while s is not None:
            obj = s.font if font else s.paragraph_format
            v = getattr(obj, attr, None)
            if v is not None:
                return v
            s = s.base_style
        return None

    # ---- body
    def add_body(self, elements, cursor):
        fresh = True  # current block is still unused
        for el in elements:
            if el.tag == qn("w:p"):
                if not fresh:
                    cursor.insertBlock()
                fresh = False
                self.add_paragraph(Paragraph(el, self.d._body), cursor)
            elif el.tag == qn("w:tbl"):
                # inserting at the end of a paragraph puts the table right after it
                self.add_table(Table(el, self.d._body), cursor)
                cursor.setBlockFormat(QTextBlockFormat())  # block after table: plain
                fresh = True

    def add_paragraph(self, p, cursor):
        style = p.style
        sname = style.name if style is not None else "Normal"
        bf = QTextBlockFormat()
        heading = 0
        m = re.match(r"Heading (\d)", sname or "")
        if m:
            heading = min(int(m.group(1)), 6)
            bf.setHeadingLevel(heading)
        elif sname == "Title":
            bf.setProperty(TITLE_PROP, True)

        def pf(attr):
            v = getattr(p.paragraph_format, attr)
            return v if v is not None else self._style_attr(style, attr)

        align = pf("alignment")
        bf.setAlignment(ALIGN_IN.get(align, Qt.AlignLeft) | Qt.AlignAbsolute)
        if pf("left_indent"):
            bf.setLeftMargin(pf("left_indent").pt * PX_PER_PT)
        if pf("right_indent"):
            bf.setRightMargin(pf("right_indent").pt * PX_PER_PT)
        if pf("first_line_indent"):
            bf.setTextIndent(pf("first_line_indent").pt * PX_PER_PT)
        before, after = pf("space_before"), pf("space_after")
        bf.setTopMargin(before.pt * PX_PER_PT if before is not None else 0)
        bf.setBottomMargin((after.pt if after is not None else self.def_after) * PX_PER_PT)
        ls, rule = pf("line_spacing"), pf("line_spacing_rule")
        if isinstance(ls, float):
            bf.setLineHeight(ls * 100, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
        elif ls is not None and rule == WD_LINE_SPACING.EXACTLY:
            bf.setLineHeight(ls.pt * PX_PER_PT, QTextBlockFormat.LineHeightTypes.FixedHeight.value)
        elif ls is not None:
            bf.setLineHeight(ls.pt * PX_PER_PT, QTextBlockFormat.LineHeightTypes.MinimumHeight.value)
        elif self.def_line:
            bf.setLineHeight(self.def_line * 100, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
        if pf("page_break_before"):
            bf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)

        numpr = self._numpr(p, style)
        if numpr:
            bf.setLeftMargin(0)
            bf.setTextIndent(0)
        cursor.setBlockFormat(bf)
        if numpr:
            self._add_to_list(cursor, *numpr)

        base = self._base_char_format(style, heading, sname == "Title")
        cursor.setBlockCharFormat(base)
        for item in p.iter_inner_content():
            if isinstance(item, Hyperlink):
                for r in item.runs:
                    self.add_run(r, base, cursor, bf, item.address or item.url)
            else:
                self.add_run(item, base, cursor, bf)

    def _numpr(self, p, style):
        def read(ppr):
            if ppr is None or ppr.numPr is None:
                return None
            n = ppr.numPr
            num_id = n.numId.val if n.numId is not None else None
            ilvl = n.ilvl.val if n.ilvl is not None else 0
            return (num_id, ilvl) if num_id else None
        r = read(p._p.pPr)
        if r is None and p._p.pPr is not None and p._p.pPr.numPr is not None:
            return None  # explicitly numbered 0 = no list
        s = style
        while r is None and s is not None:
            r = read(s.element.pPr)
            s = s.base_style
        return r

    def _num_format(self, num_id, ilvl):
        try:
            numbering = self.d.part.numbering_part.element
            abs_id = numbering.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')
            if not abs_id:
                return "bullet"
            fmt = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abs_id[0]}"]'
                                  f'/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val')
            return fmt[0] if fmt else "bullet"
        except Exception:
            return "bullet"

    def _add_to_list(self, cursor, num_id, ilvl):
        for key in [k for k in self.lists if k[0] == num_id and k[1] > ilvl]:
            del self.lists[key]  # deeper levels restart after a shallower item
        lst = self.lists.get((num_id, ilvl))
        if lst is not None:
            lst.add(cursor.block())
            return
        lf = QTextListFormat()
        lf.setStyle(LIST_FMT_IN.get(self._num_format(num_id, ilvl), QTextListFormat.ListDecimal))
        lf.setIndent(ilvl + 1)
        self.lists[(num_id, ilvl)] = cursor.createList(lf)

    def _base_char_format(self, style, heading, title):
        cf = QTextCharFormat()
        family = (self._style_font(style) if style is not None else None) or self.default_font
        size = self._style_attr(style, "size", font=True) if style is not None else None
        cf.setFontFamilies([family])
        cf.setFontPointSize(size.pt if size else self.default_size)
        if style is not None:
            for attr, setter in (("bold", lambda v: cf.setFontWeight(QFont.Bold if v else QFont.Normal)),
                                 ("italic", cf.setFontItalic)):
                v = self._style_attr(style, attr, font=True)
                if v is not None:
                    setter(v)
            s = style
            while s is not None:
                c = s.font.color
                if c is not None and c.type is not None and c.rgb is not None:
                    cf.setForeground(QColor("#" + str(c.rgb)))
                    break
                s = s.base_style
        return cf

    def add_run(self, run, base, cursor, bf, href=None):
        cf = QTextCharFormat(base)
        rpr = run._r.rPr
        if rpr is not None and rpr.find(qn("w:rStyle")) is not None and run.style is not None:
            self._apply_font(run.style.font, cf, self._style_font(run.style))
        self._apply_font(run.font, cf, self._rfonts_name(rpr.find(qn("w:rFonts")))
                         if rpr is not None else None)
        if rpr is not None:
            shd = rpr.find(qn("w:shd"))
            if shd is not None and shd.get(qn("w:fill")) not in (None, "auto"):
                cf.setBackground(QColor("#" + shd.get(qn("w:fill"))))
        if href:
            cf.setAnchor(True)
            cf.setAnchorHref(href)
            if run.font.color.rgb is None:
                cf.setForeground(QColor("#0563C1"))
            cf.setFontUnderline(True)
        buf = []

        def flush():
            if buf:
                cursor.insertText("".join(buf), cf)
                buf.clear()

        for child in run._r.iterchildren():
            tag = child.tag
            if tag == qn("w:t"):
                buf.append(child.text or "")
            elif tag == qn("w:tab"):
                buf.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                if child.get(qn("w:type")) == "page":
                    flush()
                    nbf = QTextBlockFormat(bf)
                    nbf.setPageBreakPolicy(QTextFormat.PageBreak_AlwaysBefore)
                    cursor.insertBlock(nbf, cf)
                else:
                    buf.append(" ")
            elif tag == qn("w:noBreakHyphen"):
                buf.append("‑")
            elif tag in (qn("w:drawing"), qn("w:pict"), "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
                flush()
                self._add_images(child, cursor)
        flush()

    def _apply_font(self, font, cf, family):
        if family:
            cf.setFontFamilies([family])
        if font.size is not None:
            cf.setFontPointSize(font.size.pt)
        if font.bold is not None:
            cf.setFontWeight(QFont.Bold if font.bold else QFont.Normal)
        if font.italic is not None:
            cf.setFontItalic(font.italic)
        if font.underline is not None:
            cf.setFontUnderline(bool(font.underline))
        if font.strike or font.double_strike:
            cf.setFontStrikeOut(True)
        elif font.strike is False:
            cf.setFontStrikeOut(False)
        if font.superscript:
            cf.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
        elif font.subscript:
            cf.setVerticalAlignment(QTextCharFormat.AlignSubScript)
        if font.all_caps:
            cf.setFontCapitalization(QFont.AllUppercase)
        elif font.small_caps:
            cf.setFontCapitalization(QFont.SmallCaps)
        try:
            if font.color.type is not None and font.color.rgb is not None:
                cf.setForeground(QColor("#" + str(font.color.rgb)))
            elif font.color.type == MSO_COLOR_TYPE.AUTO:
                cf.setForeground(QColor("black"))
        except Exception:
            pass
        if font.highlight_color in HIGHLIGHT:
            cf.setBackground(QColor(HIGHLIGHT[font.highlight_color]))

    def _add_images(self, el, cursor):
        blips = el.findall(".//" + qn("a:blip"))
        if not blips:
            return
        blip = blips[0]
        rid = blip.get(qn("r:embed"))
        if not rid or rid not in self.d.part.related_parts:
            return
        img = QImage.fromData(self.d.part.related_parts[rid].blob)
        if img.isNull():
            return
        ext = el.find(".//" + qn("wp:extent"))
        w, h = img.width(), img.height()
        if ext is not None:
            w, h = int(ext.get("cx")) / EMU_PER_PX, int(ext.get("cy")) / EMU_PER_PX
        self.img_count += 1
        name = f"docximg{self.img_count}.png"
        self.qdoc.addResource(QTextDocument.ImageResource, QUrl(name), img)
        f = QTextImageFormat()
        f.setName(name)
        f.setWidth(w)
        f.setHeight(h)
        cursor.insertImage(f)

    def add_table(self, tbl, cursor):
        cells = []  # (row, col, span, vmerge, tc)
        for r, tr in enumerate(tbl._tbl.tr_lst):
            col = 0
            for tc in tr.tc_lst:
                span = tc.grid_span
                cells.append((r, col, span, tc.vMerge, tc))
                col += span
        nrows = len(tbl._tbl.tr_lst)
        ncols = max([c[1] + c[2] for c in cells] + [len(tbl._tbl.tblGrid.gridCol_lst), 1])
        if not nrows:
            return
        tf = QTextTableFormat()
        tf.setBorder(1)
        tf.setBorderBrush(QBrush(QColor("#808080")))
        tf.setBorderCollapse(True)
        tf.setCellPadding(4)
        tf.setCellSpacing(0)
        widths = [gc.w for gc in tbl._tbl.tblGrid.gridCol_lst]
        if widths and all(widths) and len(widths) == ncols:
            tf.setColumnWidthConstraints([QTextLength(QTextLength.FixedLength, w.pt * PX_PER_PT)
                                          for w in widths])
        else:
            tf.setWidth(QTextLength(QTextLength.PercentageLength, 100))
        qt = cursor.insertTable(nrows, ncols, tf)
        for r, c, span, vm, tc in cells:
            if vm == "continue":
                continue
            cell = qt.cellAt(r, c)
            tcpr = tc.tcPr
            if tcpr is not None:
                shd = tcpr.find(qn("w:shd"))
                if shd is not None and shd.get(qn("w:fill")) not in (None, "auto"):
                    cfmt = cell.format().toTableCellFormat()
                    cfmt.setBackground(QColor("#" + shd.get(qn("w:fill"))))
                    cell.setFormat(cfmt)
            self.add_body([e for e in tc.iterchildren() if e.tag in (qn("w:p"), qn("w:tbl"))],
                          cell.firstCursorPosition())
        for r, c, span, vm, tc in cells:
            rows = 1
            if vm == "restart":
                while any(rc == (r + rows, c) and v == "continue"
                          for (rr, cc, _, v, _) in cells for rc in [(rr, cc)]):
                    rows += 1
            if rows > 1 or span > 1:
                qt.mergeCells(r, c, rows, span)
        cursor.setPosition(qt.lastPosition() + 1)

    # ---- page setup / header / footer
    def page_setup(self):
        ps = PageSetup()
        try:
            s = self.d.sections[0]
            if s.page_width and s.page_height:
                ps.width_in, ps.height_in = s.page_width.inches, s.page_height.inches
            for attr in ("left", "top", "right", "bottom"):
                v = getattr(s, f"{attr}_margin")
                if v is not None:
                    setattr(ps, f"margin_{attr}", v.inches)
            ps.header = "\n".join(_text_without_fields(p._p) for p in s.header.paragraphs).strip()
            ps.footer = "\n".join(_text_without_fields(p._p) for p in s.footer.paragraphs).strip()
            xml = s.footer._element.xml + s.header._element.xml
            ps.page_numbers = bool(re.search(r"\bPAGE\b", xml))
        except Exception:
            pass
        return ps


def _text_without_fields(p_el):
    out, depth = [], 0
    for el in p_el.iter():
        if el.tag == qn("w:fldChar"):
            t = el.get(qn("w:fldCharType"))
            depth += 1 if t == "begin" else -1 if t == "end" else 0
        elif el.tag == qn("w:t") and depth == 0:
            anc = el.getparent()
            while anc is not None and anc.tag != qn("w:fldSimple"):
                anc = anc.getparent()
            if anc is None:
                out.append(el.text or "")
    return "".join(out)


def load_docx(path):
    """Returns (QTextDocument, PageSetup)."""
    d = docx.Document(path)
    qdoc = QTextDocument()
    imp = _Importer(d, qdoc)
    qdoc.setDefaultFont(QFont(imp.default_font, int(round(imp.default_size))))
    cursor = QTextCursor(qdoc)
    cursor.beginEditBlock()
    imp.add_body([e for e in d.element.body.iterchildren() if e.tag in (qn("w:p"), qn("w:tbl"))],
                 cursor)
    cursor.endEditBlock()
    qdoc.setModified(False)
    return qdoc, imp.page_setup()


# ====================================================================== export

class _Exporter:
    def __init__(self, qdoc, d):
        self.qdoc, self.d = qdoc, d
        self.numbering = d.part.numbering_part.element
        self.abstract = {}    # docx format name -> abstractNumId
        self.nums = {}        # QTextList objectIndex -> numId
        self.default_family = qdoc.defaultFont().family()
        self.default_size = qdoc.defaultFont().pointSizeF()

    # ---- numbering
    def _abstract_for(self, style):
        fmt = LIST_FMT_OUT.get(style, "bullet")
        key = fmt if fmt != "bullet" else f"bullet{style.name}"
        if key in self.abstract:
            return self.abstract[key]
        ids = [int(x) for x in self.numbering.xpath("./w:abstractNum/@w:abstractNumId")]
        aid = max(ids + [0]) + 1
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(aid))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        a.append(mlt)
        for lvl in range(9):
            L = OxmlElement("w:lvl")
            L.set(qn("w:ilvl"), str(lvl))
            for tag, val in (("w:start", "1"), ("w:numFmt", fmt),
                             ("w:lvlText", BULLET_CHARS.get(style, "•") if fmt == "bullet"
                              else f"%{lvl + 1}."), ("w:lvlJc", "left")):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                L.append(e)
            ppr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 * (lvl + 1)))
            ind.set(qn("w:hanging"), "360")
            ppr.append(ind)
            L.append(ppr)
            if fmt == "bullet" and style == QTextListFormat.ListDisc:
                rpr = OxmlElement("w:rPr")
                rf = OxmlElement("w:rFonts")
                rf.set(qn("w:ascii"), "Symbol")
                rf.set(qn("w:hAnsi"), "Symbol")
                rf.set(qn("w:hint"), "default")
                rpr.append(rf)
                L.find(qn("w:lvlText")).set(qn("w:val"), "")
                L.append(rpr)
            a.append(L)
        first_num = self.numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            self.numbering.append(a)
        self.abstract[key] = aid
        return aid

    def _num_for(self, lst):
        key = lst.objectIndex()
        if key not in self.nums:
            aid = self._abstract_for(lst.format().style())
            ids = [int(x) for x in self.numbering.xpath("./w:num/@w:numId")]
            nid = max(ids + [0]) + 1
            n = OxmlElement("w:num")
            n.set(qn("w:numId"), str(nid))
            an = OxmlElement("w:abstractNumId")
            an.set(qn("w:val"), str(aid))
            n.append(an)
            self.numbering.append(n)
            self.nums[key] = nid
        return self.nums[key]

    # ---- content
    def export_frame_iter(self, it, container, first_para=None):
        while not it.atEnd():
            frame = it.currentFrame()
            if frame is not None:
                if hasattr(frame, "cellAt"):
                    self.export_table(frame, container)
                    first_para = None
            else:
                block = it.currentBlock()
                self.export_block(block, first_para or container.add_paragraph())
                first_para = None
            it += 1

    def export_block(self, block, p):
        bf = block.blockFormat()
        level = bf.headingLevel()
        if bf.property(TITLE_PROP):
            p.style = self.d.styles["Title"]
        elif 1 <= level <= 6:
            p.style = self.d.styles[f"Heading {level}"]
        pf = p.paragraph_format
        a = bf.alignment() & Qt.AlignHorizontal_Mask
        pf.alignment = (WD_ALIGN_PARAGRAPH.CENTER if a & Qt.AlignHCenter else
                        WD_ALIGN_PARAGRAPH.RIGHT if a & Qt.AlignRight else
                        WD_ALIGN_PARAGRAPH.JUSTIFY if a & Qt.AlignJustify else WD_ALIGN_PARAGRAPH.LEFT)
        pf.space_before = Pt(bf.topMargin() / PX_PER_PT)
        pf.space_after = Pt(bf.bottomMargin() / PX_PER_PT)
        lt = bf.lineHeightType()
        if lt == QTextBlockFormat.LineHeightTypes.ProportionalHeight.value:
            pf.line_spacing = bf.lineHeight() / 100
        elif lt == QTextBlockFormat.LineHeightTypes.FixedHeight.value:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(bf.lineHeight() / PX_PER_PT)
        elif lt == QTextBlockFormat.LineHeightTypes.MinimumHeight.value:
            pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            pf.line_spacing = Pt(bf.lineHeight() / PX_PER_PT)
        else:
            pf.line_spacing = 1.0
        if bf.pageBreakPolicy() & QTextFormat.PageBreak_AlwaysBefore:
            pf.page_break_before = True
        lst = block.textList()
        if lst is not None:
            numpr = p._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = max(lst.format().indent() - 1, 0)
            numpr.get_or_add_numId().val = self._num_for(lst)
        else:
            left = bf.leftMargin() + bf.indent() * self.qdoc.indentWidth()
            if left:
                pf.left_indent = Pt(left / PX_PER_PT)
            if bf.textIndent():
                pf.first_line_indent = Pt(bf.textIndent() / PX_PER_PT)
        if bf.rightMargin():
            pf.right_indent = Pt(bf.rightMargin() / PX_PER_PT)

        it = block.begin()
        while not it.atEnd():
            frag = it.fragment()
            if frag.isValid():
                self.export_fragment(p, frag)
            it += 1

    def export_fragment(self, p, frag):
        cf = frag.charFormat()
        text = frag.text()
        if cf.isImageFormat():
            imf = cf.toImageFormat()
            img = self.qdoc.resource(QTextDocument.ImageResource, QUrl(imf.name()))
            if isinstance(img, QImage):
                pass
            elif img is not None and hasattr(img, "toImage"):
                img = img.toImage()
            else:
                img = QImage(imf.name())
            if img is None or img.isNull():
                return
            w = imf.width() or img.width()
            h = imf.height() or (img.height() * w / max(img.width(), 1))
            for _ in text:
                p.add_run().add_picture(io.BytesIO(image_bytes(img)),
                                        width=Emu(int(w * EMU_PER_PX)), height=Emu(int(h * EMU_PER_PX)))
            return
        text = text.replace(" ", "\n").replace(" ", "\n").replace("￼", "")
        if not text:
            return
        if cf.isAnchor() and cf.anchorHref():
            rid = p.part.relate_to(cf.anchorHref(), RT.HYPERLINK, is_external=True)
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), rid)
            r = OxmlElement("w:r")
            h.append(r)
            p._p.append(h)
            run = Run(r, p)
            run.text = text
        else:
            run = p.add_run(text)
        f = run.font
        fams = cf.fontFamilies()
        f.name = fams[0] if fams else self.default_family
        size = cf.fontPointSize() or self.default_size
        f.size = Pt(size)
        f.bold = cf.fontWeight() >= QFont.DemiBold
        f.italic = cf.fontItalic()
        f.underline = cf.fontUnderline()
        if cf.fontStrikeOut():
            f.strike = True
        va = cf.verticalAlignment()
        if va == QTextCharFormat.AlignSuperScript:
            f.superscript = True
        elif va == QTextCharFormat.AlignSubScript:
            f.subscript = True
        cap = cf.fontCapitalization()
        if cap == QFont.AllUppercase:
            f.all_caps = True
        elif cap == QFont.SmallCaps:
            f.small_caps = True
        fg = cf.foreground()
        f.color.rgb = RGBColor.from_string(fg.color().name()[1:].upper()) \
            if fg.style() != Qt.NoBrush else RGBColor(0, 0, 0)
        bg = cf.background()
        if bg.style() != Qt.NoBrush and bg.color().alpha() > 0:
            hexc = bg.color().name().upper()
            if hexc in HIGHLIGHT_BY_HEX:
                f.highlight_color = HIGHLIGHT_BY_HEX[hexc]
            else:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), hexc[1:])
                run._r.get_or_add_rPr().append(shd)

    def export_table(self, qt, container):
        rows, cols = qt.rows(), qt.columns()
        t = container.add_table(rows=rows, cols=cols)
        try:
            t.style = self.d.styles["Table Grid"]
        except KeyError:
            pass
        cons = qt.format().columnWidthConstraints()
        if len(cons) == cols and all(c.type() == QTextLength.FixedLength for c in cons):
            for c, con in enumerate(cons):
                for r in range(rows):
                    t.cell(r, c).width = Emu(int(con.rawValue() * EMU_PER_PX))
        for r in range(rows):
            for c in range(cols):
                cell = qt.cellAt(r, c)
                if cell.row() != r or cell.column() != c:
                    continue
                dc = t.cell(r, c)
                if cell.rowSpan() > 1 or cell.columnSpan() > 1:
                    dc = dc.merge(t.cell(r + cell.rowSpan() - 1, c + cell.columnSpan() - 1))
                    for extra in dc.paragraphs[1:]:
                        extra._p.getparent().remove(extra._p)
                bg = cell.format().background()
                if bg.style() != Qt.NoBrush:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), bg.color().name()[1:].upper())
                    dc._tc.get_or_add_tcPr().append(shd)
                self.export_frame_iter(cell.begin(), dc, dc.paragraphs[0])
                if dc._tc[-1].tag == qn("w:tbl"):
                    dc.add_paragraph()


def _add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, extra in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        e = OxmlElement(tag)
        if extra:
            e.set(qn("w:fldCharType"), extra)
        else:
            e.set(qn("xml:space"), "preserve")
            e.text = f" {instr} "
        run._r.append(e)


def save_docx(qdoc, setup, path):
    d = docx.Document()
    for p in list(d.paragraphs):
        p._p.getparent().remove(p._p)
    normal = d.styles["Normal"]
    normal.font.name = qdoc.defaultFont().family()
    normal.font.size = Pt(qdoc.defaultFont().pointSizeF())
    s = d.sections[0]
    s.page_width, s.page_height = Inches(setup.width_in), Inches(setup.height_in)
    s.left_margin, s.top_margin = Inches(setup.margin_left), Inches(setup.margin_top)
    s.right_margin, s.bottom_margin = Inches(setup.margin_right), Inches(setup.margin_bottom)
    if setup.header:
        hp = s.header.paragraphs[0]
        hp.text = setup.header
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if setup.footer or setup.page_numbers:
        fp = s.footer.paragraphs[0]
        fp.text = setup.footer
        if setup.page_numbers:
            np_ = s.footer.add_paragraph() if setup.footer else fp
            np_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _add_field(np_, "PAGE")
    ex = _Exporter(qdoc, d)
    ex.export_frame_iter(qdoc.rootFrame().begin(), d)
    d.save(path)
