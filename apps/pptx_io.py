"""Presentation model + .pptx import/export for Ethical Slides (uses python-pptx + lxml).

Coordinates are pixels at 96 dpi (1 px = 9525 EMU). Text lives in QTextDocuments so the editor,
renderer and exporter all share one representation.
"""
import copy
import io
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation as PptxPresentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt
from PySide6.QtCore import Qt
from PySide6.QtGui import (QBrush, QColor, QFont, QTextBlockFormat, QTextCharFormat, QTextCursor,
                           QTextDocument, QTextLength, QTextListFormat, QTextTableFormat)

EMU = 9525
PX_PER_PT = 96 / 72
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
INDENT_PX = 24


def q(tag):
    pre, name = tag.split(":")
    return "{%s}%s" % (NS[pre], name)


def find(el, path):
    return el.find(path, NS) if el is not None else None


def findall(el, path):
    return el.findall(path, NS) if el is not None else []


# ====================================================================== model

def clone_doc(doc):
    """QTextDocument.clone() drops list membership of empty paragraphs - restore it."""
    new = doc.clone()
    new.setIndentWidth(doc.indentWidth())
    lists = {}
    a, b = doc.begin(), new.begin()
    while a.isValid() and b.isValid():
        lst = a.textList()
        if lst is not None and b.textList() is None:
            key = lst.objectIndex()
            if key in lists:
                lists[key].add(b)
            else:
                lists[key] = QTextCursor(b).createList(lst.format())
        elif lst is not None:
            lists.setdefault(lst.objectIndex(), b.textList())
        a, b = a.next(), b.next()
    new.setModified(False)
    return new


@dataclass
class Element:
    kind: str = "shape"                 # shape | image | line | table
    geom: str = "rect"                  # rect roundrect ellipse triangle ... | custom
    x: float = 0
    y: float = 0
    w: float = 100
    h: float = 50
    rot: float = 0
    flip_h: bool = False
    flip_v: bool = False
    fill: object = None                 # "#RRGGBB" or None
    line: object = None
    line_w: float = 1.33
    doc: object = None                  # QTextDocument or None
    anchor: str = "t"                   # t | m | b
    insets: tuple = (9.6, 4.8, 9.6, 4.8)
    wrap: bool = True
    image: bytes = None
    crop: tuple = (0, 0, 0, 0)          # fractions l, t, r, b
    arrow_start: bool = False
    arrow_end: bool = False
    role: str = "other"                 # title | body | other
    prompt: str = ""                    # grey hint shown while empty (editor only)
    locked: bool = False                # master/layout decoration
    cust_xml: str = None                # a:custGeom xml for freeform shapes

    def clone(self):
        e = copy.copy(self)
        if self.doc is not None:
            e.doc = clone_doc(self.doc)
        return e

    def is_empty(self):
        return self.doc is None or not self.doc.toPlainText().strip()


@dataclass
class Layout:
    name: str
    index: int = None                   # layout index in the template file
    decor: list = field(default_factory=list)
    bg: object = None
    placeholders: list = field(default_factory=list)   # Element templates


@dataclass
class Slide:
    elements: list = field(default_factory=list)
    bg: object = None                   # ("color", hex) | ("image", bytes) | None (white)
    bg_override: bool = True            # write background explicitly when saving
    notes: str = ""
    layout: int = None                  # index into Presentation.layouts
    decor: list = field(default_factory=list)
    hidden: bool = False

    def clone(self):
        s = copy.copy(self)
        s.elements = [e.clone() for e in self.elements]
        s.decor = list(self.decor)      # decorations are read-only, share them
        return s


@dataclass
class Presentation:
    width: float = 1280
    height: float = 720
    slides: list = field(default_factory=list)
    layouts: list = field(default_factory=list)
    template: bytes = None              # the original .pptx file (keeps its masters/themes/layouts)
    theme: dict = None


# ====================================================================== themes

class ThemeCtx:
    def __init__(self, master):
        self.colors, self.major, self.minor = {}, "Calibri Light", "Calibri"
        try:
            theme_part = master.part.part_related_by(RT.THEME)
            root = etree.fromstring(theme_part.blob)
            scheme = find(root, ".//a:clrScheme")
            for child in scheme:
                name = etree.QName(child).localname
                c = child[0]
                self.colors[name] = c.get("val") if etree.QName(c).localname == "srgbClr" else c.get("lastClr", "000000")
            ma, mi = find(root, ".//a:majorFont/a:latin"), find(root, ".//a:minorFont/a:latin")
            self.major = ma.get("typeface") if ma is not None else self.major
            self.minor = mi.get("typeface") if mi is not None else self.minor
        except Exception:
            pass
        defaults = {"dk1": "000000", "lt1": "FFFFFF", "dk2": "44546A", "lt2": "E7E6E6", "accent1": "4472C4",
                    "accent2": "ED7D31", "accent3": "A5A5A5", "accent4": "FFC000", "accent5": "5B9BD5",
                    "accent6": "70AD47", "hlink": "0563C1", "folHlink": "954F72"}
        for k, v in defaults.items():
            self.colors.setdefault(k, v)
        cm = find(master._element, "p:clrMap")
        self.map = dict(cm.attrib) if cm is not None else {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}

    def font(self, typeface):
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.major
        if typeface.startswith("+mn"):
            return self.minor
        return typeface

    def color(self, parent, ph_color=None):
        """Resolve the colour child of parent (solidFill, fillRef, ...) to '#RRGGBB'."""
        if parent is None:
            return None
        for c in parent:
            tag = etree.QName(c).localname
            base = None
            if tag == "srgbClr":
                base = c.get("val")
            elif tag == "schemeClr":
                v = c.get("val")
                if v == "phClr":
                    base = (ph_color or "#000000").lstrip("#")
                else:
                    base = self.colors.get(self.map.get(v, v), self.colors.get(v))
            elif tag == "sysClr":
                base = c.get("lastClr", "000000")
            elif tag == "prstClr":
                base = QColor(c.get("val", "black")).name()[1:]
            elif tag == "scrgbClr":
                base = "%02X%02X%02X" % tuple(int(int(c.get(k, 0)) / 100000 * 255) for k in ("r", "g", "b"))
            if base is None:
                continue
            col = QColor("#" + base)
            for mod in c:
                m, val = etree.QName(mod).localname, int(mod.get("val", 0)) / 100000
                if m == "lumMod":
                    h, s, l, a = col.getHslF()
                    col = QColor.fromHslF(max(h, 0), s, min(1, l * val), a)
                elif m == "lumOff":
                    h, s, l, a = col.getHslF()
                    col = QColor.fromHslF(max(h, 0), s, min(1, l + val), a)
                elif m == "tint":
                    col = QColor(*(int(v + (255 - v) * (1 - val)) for v in (col.red(), col.green(), col.blue())))
                elif m == "shade":
                    col = QColor(*(int(v * val) for v in (col.red(), col.green(), col.blue())))
            return col.name().upper()
        return None


# ====================================================================== import helpers

UNSET = object()
PRESETS = {"rect": "rect", "roundRect": "roundrect", "ellipse": "ellipse", "triangle": "triangle",
           "rtTriangle": "rttriangle", "diamond": "diamond", "rightArrow": "rarrow", "leftArrow": "larrow",
           "upArrow": "uarrow", "downArrow": "darrow", "pentagon": "pentagon", "hexagon": "hexagon",
           "homePlate": "homeplate", "chevron": "chevron", "star5": "star5", "parallelogram": "parallelogram",
           "trapezoid": "trapezoid", "line": "line", "straightConnector1": "line", "flowChartProcess": "rect",
           "snip1Rect": "rect", "round2SameRect": "roundrect", "plaque": "roundrect", "frame": "rect",
           "flowChartAlternateProcess": "roundrect", "octagon": "octagon", "cloud": "ellipse",
           "heart": "ellipse", "donut": "ellipse", "flowChartConnector": "ellipse", "plus": "plus"}


def _fill(spPr, theme, style=None):
    """spPr fill -> '#hex', None (no fill) or UNSET (inherit)."""
    if spPr is not None:
        if find(spPr, "a:noFill") is not None:
            return None
        sf = find(spPr, "a:solidFill")
        if sf is not None:
            return theme.color(sf)
        gf = find(spPr, "a:gradFill")
        if gf is not None:
            stops = findall(gf, "a:gsLst/a:gs")
            if stops:
                return theme.color(stops[len(stops) // 2])
        if find(spPr, "a:blipFill") is not None or find(spPr, "a:pattFill") is not None:
            return UNSET
    if style is not None:
        ref = find(style, "a:fillRef")
        if ref is not None and ref.get("idx", "0") != "0":
            return theme.color(ref)
    return UNSET


def _line(spPr, theme, style=None):
    """-> (colour or None or UNSET, width px or None)."""
    ln = find(spPr, "a:ln")
    width = int(ln.get("w")) / EMU if ln is not None and ln.get("w") else None
    if ln is not None:
        if find(ln, "a:noFill") is not None:
            return None, width
        sf = find(ln, "a:solidFill")
        if sf is not None:
            return theme.color(sf), width
    if style is not None:
        ref = find(style, "a:lnRef")
        if ref is not None and ref.get("idx", "0") != "0":
            return theme.color(ref), width or 1.33
    return UNSET, width


def _background(cSld, theme, part):
    bg = find(cSld, "p:bg")
    if bg is None:
        return None
    pr = find(bg, "p:bgPr")
    if pr is not None:
        blip = find(pr, "a:blipFill/a:blip")
        if blip is not None:
            rid = blip.get(q("r:embed"))
            try:
                return ("image", part.related_part(rid).blob)
            except Exception:
                return None
        c = _fill(pr, theme)
        if c is not UNSET and c:
            return ("color", c)
        return None
    ref = find(bg, "p:bgRef")
    if ref is not None:
        c = theme.color(ref)
        if c:
            return ("color", c)
    return None


class TextStyle:
    """Lookup chain for paragraph/run properties: own pPr, then lstStyles from shape/layout/master."""

    def __init__(self, lst_styles, theme, default_size=18, default_color="#000000"):
        self.lst = [ls for ls in lst_styles if ls is not None]
        self.theme = theme
        self.default_size, self.default_color = default_size, default_color

    def ppr_chain(self, pPr, level):
        chain = [pPr] if pPr is not None else []
        for ls in self.lst:
            lv = find(ls, f"a:lvl{level + 1}pPr")
            if lv is not None:
                chain.append(lv)
        return chain

    @staticmethod
    def attr(chain, name):
        for el in chain:
            if el is not None and el.get(name) is not None:
                return el.get(name)
        return None

    def bullet(self, chain):
        for el in chain:
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip"):
                b = find(el, tag)
                if b is not None:
                    if tag == "a:buNone":
                        return None
                    if tag == "a:buAutoNum":
                        return ("num", b.get("type", "arabicPeriod"))
                    return ("char", b.get("char", "•"))
        return None

    def spacing(self, chain, tag):
        for el in chain:
            sp = find(el, tag)
            if sp is not None:
                pts = find(sp, "a:spcPts")
                if pts is not None:
                    return ("pt", int(pts.get("val")) / 100)
                pct = find(sp, "a:spcPct")
                if pct is not None:
                    return ("pct", int(pct.get("val")) / 1000)
        return None

    def run_format(self, rPr, chain, scale=1.0, ph_color=None):
        rchain = ([rPr] if rPr is not None else []) + [find(el, "a:defRPr") for el in chain]
        rchain = [r for r in rchain if r is not None]
        cf = QTextCharFormat()
        sz = self.attr(rchain, "sz")
        cf.setFontPointSize(max(1.0, (int(sz) / 100 if sz else self.default_size) * scale))
        b, i, u = self.attr(rchain, "b"), self.attr(rchain, "i"), self.attr(rchain, "u")
        cf.setFontWeight(QFont.Bold if b in ("1", "true") else QFont.Normal)
        cf.setFontItalic(i in ("1", "true"))
        cf.setFontUnderline(u not in (None, "none"))
        if self.attr(rchain, "strike") in ("sngStrike", "dblStrike"):
            cf.setFontStrikeOut(True)
        cap = self.attr(rchain, "cap")
        if cap == "all":
            cf.setFontCapitalization(QFont.AllUppercase)
        elif cap == "small":
            cf.setFontCapitalization(QFont.SmallCaps)
        base = self.attr(rchain, "baseline")
        if base and int(base) > 0:
            cf.setVerticalAlignment(QTextCharFormat.AlignSuperScript)
        elif base and int(base) < 0:
            cf.setVerticalAlignment(QTextCharFormat.AlignSubScript)
        family, color = None, None
        if ph_color is not None:  # a shape's own text colour beats the master's default
            own = find(rPr, "a:solidFill") if rPr is not None else None
            color = self.theme.color(own) if own is not None else ph_color
        for r in rchain:
            if family is None:
                lat = find(r, "a:latin")
                if lat is not None:
                    family = self.theme.font(lat.get("typeface"))
            if color is None:
                sf = find(r, "a:solidFill")
                if sf is not None:
                    color = self.theme.color(sf)
        cf.setFontFamilies([family or self.theme.minor])
        cf.setForeground(QColor(color or self.default_color))
        return cf


def build_doc(txBody, style, scale=1.0, ph_color=None):
    """a:txBody -> QTextDocument."""
    doc = QTextDocument()
    doc.setDocumentMargin(0)
    doc.setIndentWidth(INDENT_PX)
    doc.setDefaultFont(QFont(style.theme.minor, int(style.default_size)))
    cur = QTextCursor(doc)
    lists = {}
    paras = findall(txBody, "a:p") or [None]
    for n, p in enumerate(paras):
        pPr = find(p, "a:pPr")
        level = int(pPr.get("lvl", 0)) if pPr is not None else 0
        chain = style.ppr_chain(pPr, level)
        bf = QTextBlockFormat()
        algn = style.attr(chain, "algn")
        bf.setAlignment({"ctr": Qt.AlignHCenter, "r": Qt.AlignRight, "just": Qt.AlignJustify,
                         "dist": Qt.AlignJustify}.get(algn, Qt.AlignLeft) | Qt.AlignAbsolute)
        marL = style.attr(chain, "marL")
        indent = style.attr(chain, "indent")
        bullet = style.bullet(chain)
        marL_px = int(marL) / EMU if marL else 0
        if bullet:
            bf.setLeftMargin(max(0, marL_px - (level + 1) * INDENT_PX))
        else:
            bf.setLeftMargin(marL_px)
            if indent:
                bf.setTextIndent(int(indent) / EMU)
        for tag, setter in (("a:spcBef", bf.setTopMargin), ("a:spcAft", bf.setBottomMargin)):
            sp = style.spacing(chain, tag)
            if sp:
                size = style.default_size
                setter(sp[1] * PX_PER_PT if sp[0] == "pt" else sp[1] / 100 * size * 1.2 * PX_PER_PT * scale)
        ls = style.spacing(chain, "a:lnSpc")
        if ls and ls[0] == "pct":
            bf.setLineHeight(ls[1] * 0.9 if ls[1] > 0 else 100, QTextBlockFormat.LineHeightTypes.ProportionalHeight.value)
        elif ls:
            bf.setLineHeight(ls[1] * PX_PER_PT, QTextBlockFormat.LineHeightTypes.FixedHeight.value)
        if n:
            cur.insertBlock(bf)
        else:
            cur.setBlockFormat(bf)
        end = find(p, "a:endParaRPr")
        cur.setBlockCharFormat(style.run_format(end, chain, scale, ph_color))
        if bullet:
            key = (level, bullet[0])
            lst = lists.get(key)
            if lst is None:
                lf = QTextListFormat()
                if bullet[0] == "num":
                    t = bullet[1]
                    lf.setStyle(QTextListFormat.ListLowerAlpha if t.startswith("alphaLc") else
                                QTextListFormat.ListUpperAlpha if t.startswith("alphaUc") else
                                QTextListFormat.ListLowerRoman if t.startswith("romanLc") else
                                QTextListFormat.ListUpperRoman if t.startswith("romanUc") else QTextListFormat.ListDecimal)
                else:
                    lf.setStyle(QTextListFormat.ListCircle if bullet[1] in ("o", "○", "◦") else
                                QTextListFormat.ListSquare if bullet[1] in ("▪", "■", "", "§") else
                                QTextListFormat.ListDisc)
                lf.setIndent(level + 1)
                lists[key] = cur.createList(lf)
            else:
                lst.add(cur.block())
        else:
            lists.clear()
        if p is None:
            continue
        for child in p:
            tag = etree.QName(child).localname
            if tag in ("r", "fld"):
                t = find(child, "a:t")
                cur.insertText(t.text if t is not None and t.text else "",
                               style.run_format(find(child, "a:rPr"), chain, scale, ph_color))
            elif tag == "br":
                cur.insertText(" ", style.run_format(find(child, "a:rPr"), chain, scale, ph_color))
    doc.setModified(False)
    return doc


def _geom_from(spPr):
    pg = find(spPr, "a:prstGeom")
    if pg is not None:
        return PRESETS.get(pg.get("prst"), "rect"), None
    cg = find(spPr, "a:custGeom")
    if cg is not None:
        return "custom", etree.tostring(cg).decode()
    return None, None


def _xfrm(el, tag="a:xfrm"):
    x = find(el, tag)
    if x is None:
        return None
    off, ext = find(x, "a:off"), find(x, "a:ext")
    if off is None or ext is None:
        return None
    return dict(x=int(off.get("x")) / EMU, y=int(off.get("y")) / EMU, w=int(ext.get("cx")) / EMU,
                h=int(ext.get("cy")) / EMU, rot=int(x.get("rot", 0)) / 60000,
                flip_h=x.get("flipH") in ("1", "true"), flip_v=x.get("flipV") in ("1", "true"))


class _Importer:
    def __init__(self, prs):
        self.prs = prs
        self.warnings = set()
        self.themes = {}

    def theme(self, master):
        key = id(master.part)
        if key not in self.themes:
            self.themes[key] = ThemeCtx(master)
        return self.themes[key]

    # ---- placeholders: find the layout + master shapes this one inherits from
    @staticmethod
    def _ph_info(sp):
        ph = find(sp, "p:nvSpPr/p:nvPr/p:ph")
        if ph is None:
            ph = find(sp, "p:nvPicPr/p:nvPr/p:ph")
        if ph is None:
            return None
        return ph.get("type", "body" if ph.get("idx") else "body"), ph.get("idx")

    @staticmethod
    def _match_ph(container, ph_type, idx):
        best = None
        for sp in findall(container, "p:cSld/p:spTree/p:sp"):
            ph = find(sp, "p:nvSpPr/p:nvPr/p:ph")
            if ph is None:
                continue
            t = ph.get("type", "body")
            if idx is not None and ph.get("idx") == idx and t == ph_type:
                return sp
            if idx is not None and ph.get("idx") == idx and best is None:
                best = sp
            if t == ph_type or (ph_type in ("title", "ctrTitle") and t in ("title", "ctrTitle")):
                best = best if best is not None else sp
        return best

    def convert_tree(self, spTree, part, theme, masters, ctx_kind, out, xf=(1, 1, 0, 0), skip_ph=False):
        """Convert children of an spTree (or grpSp) into Elements appended to out.
        masters: [(layout_el, ...), (master_el, ...)] for placeholder inheritance."""
        sx, sy, tx, ty = xf
        for el in spTree:
            tag = etree.QName(el).localname
            try:
                if tag == "grpSp":
                    gx = find(el, "p:grpSpPr/a:xfrm")
                    if gx is None:
                        self.convert_tree(el, part, theme, masters, ctx_kind, out, xf, skip_ph)
                        continue
                    off, ext = find(gx, "a:off"), find(gx, "a:ext")
                    choff, chext = find(gx, "a:chOff"), find(gx, "a:chExt")
                    ox, oy = int(off.get("x")) / EMU, int(off.get("y")) / EMU
                    ex, ey = int(ext.get("cx")) / EMU, int(ext.get("cy")) / EMU
                    cox, coy = int(choff.get("x")) / EMU, int(choff.get("y")) / EMU
                    cex, cey = max(int(chext.get("cx")) / EMU, 1e-6), max(int(chext.get("cy")) / EMU, 1e-6)
                    nsx, nsy = ex / cex, ey / cey
                    # child -> group: x' = ox + (x - cox) * nsx ; then apply outer transform
                    inner = (sx * nsx, sy * nsy, tx + sx * (ox - cox * nsx), ty + sy * (oy - coy * nsy))
                    self.convert_tree(el, part, theme, masters, ctx_kind, out, inner, skip_ph)
                elif tag in ("sp", "cxnSp"):
                    e = self.convert_sp(el, part, theme, masters, ctx_kind, skip_ph)
                    if e:
                        self._apply_xf(e, xf)
                        out.append(e)
                elif tag == "pic":
                    e = self.convert_pic(el, part, masters, skip_ph)
                    if e:
                        self._apply_xf(e, xf)
                        out.append(e)
                elif tag == "graphicFrame":
                    e = self.convert_frame(el, theme)
                    if e:
                        self._apply_xf(e, xf)
                        out.append(e)
                elif tag == "AlternateContent":
                    fb = el.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback")
                    if fb is not None:
                        self.convert_tree(fb, part, theme, masters, ctx_kind, out, xf, skip_ph)
            except Exception:
                self.warnings.add("some shapes")

    @staticmethod
    def _apply_xf(e, xf):
        sx, sy, tx, ty = xf
        if (sx, sy, tx, ty) == (1, 1, 0, 0):
            return
        e.x, e.y = tx + e.x * sx, ty + e.y * sy
        e.w, e.h = e.w * sx, e.h * sy

    def _inherited(self, sp, masters):
        """[shape, layout ph, master ph] chain for a placeholder (missing ones skipped)."""
        info = self._ph_info(sp)
        chain = [sp]
        if info:
            for m in masters:
                match = self._match_ph(m, *info)
                if match is not None:
                    chain.append(match)
        return chain, info

    def convert_sp(self, sp, part, theme, masters, ctx_kind, skip_ph):
        chain, info = self._inherited(sp, masters)
        if info and skip_ph:
            return None
        if info and info[0] in ("dt", "ftr", "sldNum") and ctx_kind != "slide":
            return None
        e = Element()
        is_cxn = etree.QName(sp).localname == "cxnSp"
        xf = None
        for s in chain:
            xf = _xfrm(find(s, "p:spPr"))
            if xf:
                break
        if not xf:
            return None
        for k, v in xf.items():
            setattr(e, k, v)
        spPr = find(sp, "p:spPr")
        style = find(sp, "p:style")
        geom, cust = None, None
        for s in chain:
            geom, cust = _geom_from(find(s, "p:spPr"))
            if geom:
                break
        e.geom = geom or "rect"
        e.cust_xml = cust
        if is_cxn or e.geom == "line":
            e.kind, e.geom = "line", "line"
        fill = UNSET
        for s in chain:
            fill = _fill(find(s, "p:spPr"), theme, find(s, "p:style"))
            if fill is not UNSET:
                break
        e.fill = None if fill is UNSET else fill
        line, lw = UNSET, None
        for s in chain:
            line, w = _line(find(s, "p:spPr"), theme, find(s, "p:style"))
            lw = lw or w
            if line is not UNSET:
                break
        e.line = None if line is UNSET else line
        e.line_w = lw or 1.33
        ln = find(spPr, "a:ln")
        if ln is not None:
            te, he = find(ln, "a:tailEnd"), find(ln, "a:headEnd")
            e.arrow_end = te is not None and te.get("type", "none") != "none"
            e.arrow_start = he is not None and he.get("type", "none") != "none"
        if e.kind == "line":
            if e.line is None and line is UNSET:
                e.line = "#000000"
            return e
        # text
        ph_type = info[0] if info else None
        e.role = "title" if ph_type in ("title", "ctrTitle") else "body" if info and ph_type not in ("dt", "ftr", "sldNum") else "other"
        cnv = find(sp, "p:nvSpPr/p:cNvPr")
        if not info and cnv is not None and cnv.get("name", "").startswith("Ethical "):
            e.role = {"Ethical Title": "title", "Ethical Body": "body"}.get(cnv.get("name"), "other")
        txBody = find(sp, "p:txBody")
        bodies = [find(s, "p:txBody") for s in chain]
        lst = [find(b, "a:lstStyle") for b in bodies]
        master_el = masters[-1] if masters else None
        tx_styles = find(master_el, "p:txStyles")
        if ph_type in ("title", "ctrTitle"):
            lst.append(find(tx_styles, "p:titleStyle"))
            dsize = 44
        elif info and ph_type not in ("dt", "ftr", "sldNum"):
            lst.append(find(tx_styles, "p:bodyStyle"))
            dsize = 28
        else:
            lst.append(find(tx_styles, "p:otherStyle"))
            dsize = 18
        font_ref = find(style, "a:fontRef")
        ph_color = theme.color(font_ref) if font_ref is not None else None
        tstyle = TextStyle(lst, theme, dsize, "#" + theme.colors.get(theme.map.get("tx1", "dk1"), "000000"))
        bodyPr = None
        for b in bodies:
            bp = find(b, "a:bodyPr")
            if bp is not None:
                bodyPr = bp if bodyPr is None else bodyPr
        anchor = TextStyle.attr([find(b, "a:bodyPr") for b in bodies], "anchor")
        e.anchor = {"ctr": "m", "b": "b"}.get(anchor, "t")
        ins = [TextStyle.attr([find(b, "a:bodyPr") for b in bodies], k) for k in ("lIns", "tIns", "rIns", "bIns")]
        e.insets = tuple(int(v) / EMU if v is not None else d for v, d in zip(ins, (9.6, 4.8, 9.6, 4.8)))
        wrap = TextStyle.attr([find(b, "a:bodyPr") for b in bodies], "wrap")
        e.wrap = wrap != "none"
        scale = 1.0
        na = find(find(txBody, "a:bodyPr"), "a:normAutofit")
        if na is not None and na.get("fontScale"):
            scale = int(na.get("fontScale")) / 100000
        has_text = txBody is not None and any((t.text or "").strip() for t in findall(txBody, ".//a:t"))
        if txBody is not None and (has_text or info):
            e.doc = build_doc(txBody if has_text else self._prompt_body(bodies), tstyle, scale, ph_color)
            if not has_text:
                e.doc = self._emptied(e.doc)
                e.prompt = self._prompt_text(bodies, ph_type)
        if ctx_kind != "slide" and info:
            return None
        return e

    @staticmethod
    def _prompt_body(bodies):
        for b in bodies:
            if b is not None and findall(b, "a:p"):
                return b
        return None

    @staticmethod
    def _emptied(doc):
        """Keep the first block's formatting but drop the text."""
        c = QTextCursor(doc)
        c.movePosition(QTextCursor.Start)
        fmt = c.block().charFormat()
        bf = c.blockFormat()
        lst = c.block().textList()
        lf = lst.format() if lst else None
        new = QTextDocument()
        new.setDocumentMargin(0)
        new.setIndentWidth(INDENT_PX)
        new.setDefaultFont(doc.defaultFont())
        nc = QTextCursor(new)
        nc.setBlockFormat(bf)
        nc.setBlockCharFormat(fmt)
        nc.setCharFormat(fmt)
        if lf is not None:
            nc.createList(lf)
        new.setModified(False)
        return new

    @staticmethod
    def _prompt_text(bodies, ph_type):
        if ph_type in ("title", "ctrTitle"):
            return "Click to add title"
        if ph_type == "subTitle":
            return "Click to add subtitle"
        return "Click to add text"

    def convert_pic(self, pic, part, masters, skip_ph):
        chain, info = self._inherited(pic, masters) if find(pic, "p:nvPicPr/p:nvPr/p:ph") is not None else ([pic], None)
        if info and skip_ph:
            return None
        xf = None
        for s in chain:
            xf = _xfrm(find(s, "p:spPr"))
            if xf:
                break
        blip = find(pic, "p:blipFill/a:blip")
        if xf is None or blip is None:
            return None
        rid = blip.get(q("r:embed"))
        try:
            data = part.related_part(rid).blob
        except Exception:
            return None
        e = Element(kind="image", image=data)
        for k, v in xf.items():
            setattr(e, k, v)
        src = find(pic, "p:blipFill/a:srcRect")
        if src is not None:
            e.crop = tuple(int(src.get(k, 0)) / 100000 for k in ("l", "t", "r", "b"))
        if find(pic, "p:nvPicPr/p:nvPr/a:videoFile") is not None or find(pic, "p:nvPicPr/p:nvPr/a:audioFile") is not None:
            self.warnings.add("video/audio")
        return e

    def convert_frame(self, gf, theme):
        xf = _xfrm(gf, "p:xfrm")
        if xf is None:
            return None
        data = find(gf, "a:graphic/a:graphicData")
        uri = data.get("uri", "") if data is not None else ""
        tbl = find(data, "a:tbl")
        if tbl is None:
            self.warnings.add("charts" if "chart" in uri else "SmartArt" if "diagram" in uri else "embedded objects")
            return None
        e = Element(kind="table")
        for k, v in xf.items():
            setattr(e, k, v)
        e.doc = self.table_doc(tbl, theme)
        e.wrap = True
        e.insets = (0, 0, 0, 0)
        return e

    def table_doc(self, tbl, theme):
        doc = QTextDocument()
        doc.setDocumentMargin(0)
        doc.setIndentWidth(INDENT_PX)
        doc.setDefaultFont(QFont(theme.minor, 18))
        cols = [int(g.get("w")) / EMU for g in findall(tbl, "a:tblGrid/a:gridCol")]
        rows = findall(tbl, "a:tr")
        tf = QTextTableFormat()
        tf.setBorder(1)
        tf.setBorderBrush(QBrush(QColor("#FFFFFF")))
        tf.setBorderCollapse(True)
        tf.setCellPadding(6)
        tf.setCellSpacing(0)
        tf.setColumnWidthConstraints([QTextLength(QTextLength.FixedLength, w) for w in cols])
        cur = QTextCursor(doc)
        t = cur.insertTable(max(len(rows), 1), max(len(cols), 1), tf)
        tblPr = find(tbl, "a:tblPr")
        first_row = tblPr is not None and tblPr.get("firstRow") in ("1", "true")
        banded = tblPr is not None and tblPr.get("bandRow") in ("1", "true")
        accent = "#" + theme.colors.get("accent1", "4472C4")
        light1 = QColor(accent).lighter(170).name()
        light2 = QColor(accent).lighter(185).name()
        tstyle = TextStyle([], theme, 18, "#000000")
        spans = []
        for r, tr in enumerate(rows):
            for c, tc in enumerate(findall(tr, "a:tc")):
                if c >= len(cols):
                    break
                if tc.get("hMerge") or tc.get("vMerge"):
                    continue
                gs, rs = int(tc.get("gridSpan", 1)), int(tc.get("rowSpan", 1))
                if gs > 1 or rs > 1:
                    spans.append((r, c, rs, gs))
                cell = t.cellAt(r, c)
                tcPr = find(tc, "a:tcPr")
                fill = _fill(tcPr, theme) if tcPr is not None else UNSET
                header = first_row and r == 0
                if fill is UNSET:
                    fill = accent if header else (light1 if banded and (r - first_row) % 2 == 0 else light2) \
                        if (first_row or banded) else None
                fmt = cell.format().toTableCellFormat()
                if fill:
                    fmt.setBackground(QColor(fill))
                cell.setFormat(fmt)
                cdoc = build_doc(find(tc, "a:txBody"), TextStyle([], theme, 18, "#FFFFFF" if header else "#000000"),
                                 ph_color=None)
                cc = cell.firstCursorPosition()
                self._copy_doc_into(cdoc, cc, bold=header)
        for r, c, rs, gs in spans:
            t.mergeCells(r, c, rs, gs)
        doc.setModified(False)
        return doc

    @staticmethod
    def _copy_doc_into(src, cursor, bold=False):
        b = src.begin()
        first = True
        while b.isValid():
            if not first:
                cursor.insertBlock(b.blockFormat(), b.charFormat())
            else:
                cursor.setBlockFormat(b.blockFormat())
            first = False
            it = b.begin()
            while not it.atEnd():
                frag = it.fragment()
                if frag.isValid():
                    cf = frag.charFormat()
                    if bold:
                        cf.setFontWeight(QFont.Bold)
                    cursor.insertText(frag.text(), cf)
                it += 1
            b = b.next()

    # ---- layouts / masters
    def layout_info(self, layout, index):
        master = layout.slide_master
        theme = self.theme(master)
        decor = []
        if find(master._element, "p:cSld") is not None and layout._element.get("showMasterSp", "1") not in ("0", "false"):
            self.convert_tree(find(master._element, "p:cSld/p:spTree"), master.part, theme,
                              [master._element], "master", decor, skip_ph=True)
        self.convert_tree(find(layout._element, "p:cSld/p:spTree"), layout.part, theme,
                          [layout._element, master._element], "layout", decor, skip_ph=True)
        for d in decor:
            d.locked = True
        bg = _background(find(layout._element, "p:cSld"), theme, layout.part) or \
            _background(find(master._element, "p:cSld"), theme, master.part)
        phs = []
        for sp in findall(layout._element, "p:cSld/p:spTree/p:sp"):
            info = self._ph_info(sp)
            if not info or info[0] in ("dt", "ftr", "sldNum"):
                continue
            e = self.convert_sp(sp, layout.part, theme, [master._element], "slide", False)
            if e is not None:
                if e.doc is not None and not e.prompt:
                    e.doc = self._emptied(e.doc)
                    e.prompt = self._prompt_text(None, info[0])
                e.fill = e.fill if e.fill else None
                phs.append(e)
        return Layout(layout.name, index, decor, bg, phs)

    def slide(self, s, layouts, layout_index):
        layout = s.slide_layout
        master = layout.slide_master
        theme = self.theme(master)
        out = Slide()
        out.layout = layout_index
        lay = layouts[layout_index]
        show_master = s._element.get("showMasterSp", "1") not in ("0", "false")
        out.decor = list(lay.decor) if show_master else []
        own_bg = _background(find(s._element, "p:cSld"), theme, s.part)
        out.bg = own_bg or lay.bg
        out.bg_override = own_bg is not None
        self.convert_tree(find(s._element, "p:cSld/p:spTree"), s.part, theme,
                          [layout._element, master._element], "slide", out.elements)
        out.hidden = s._element.get("show") in ("0", "false")
        if s.has_notes_slide:
            try:
                out.notes = s.notes_slide.notes_text_frame.text
            except Exception:
                pass
        if find(s._element, "p:transition") is not None or find(s._element, "p:timing") is not None:
            self.warnings.add("animations/transitions")
        return out


def load_pptx(path):
    """Returns (Presentation, sorted list of warnings)."""
    prs = PptxPresentation(path)
    imp = _Importer(prs)
    with open(path, "rb") as fh:
        raw = fh.read()
    pres = Presentation(width=prs.slide_width / EMU, height=prs.slide_height / EMU, template=raw)
    all_layouts = [lay for m in prs.slide_masters for lay in m.slide_layouts]
    pres.layouts = [imp.layout_info(lay, i) for i, lay in enumerate(all_layouts)]
    index_of = {id(lay.part): i for i, lay in enumerate(all_layouts)}
    for s in prs.slides:
        pres.slides.append(imp.slide(s, pres.layouts, index_of.get(id(s.slide_layout.part), 0)))
    if not pres.slides:
        pres.slides.append(new_slide_from_layout(pres, 0))
    return pres, sorted(imp.warnings)


def new_slide_from_layout(pres, li):
    lay = pres.layouts[li]
    s = Slide(layout=li, decor=list(lay.decor), bg=lay.bg, bg_override=pres.template is None)
    s.elements = [e.clone() for e in lay.placeholders]
    return s


# ====================================================================== export

SHAPE_OUT = {"rect": MSO_SHAPE.RECTANGLE, "roundrect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
             "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "rttriangle": MSO_SHAPE.RIGHT_TRIANGLE,
             "diamond": MSO_SHAPE.DIAMOND, "rarrow": MSO_SHAPE.RIGHT_ARROW, "larrow": MSO_SHAPE.LEFT_ARROW,
             "uarrow": MSO_SHAPE.UP_ARROW, "darrow": MSO_SHAPE.DOWN_ARROW, "pentagon": MSO_SHAPE.REGULAR_PENTAGON,
             "hexagon": MSO_SHAPE.HEXAGON, "homeplate": MSO_SHAPE.PENTAGON, "chevron": MSO_SHAPE.CHEVRON,
             "star5": MSO_SHAPE.STAR_5_POINT, "parallelogram": MSO_SHAPE.PARALLELOGRAM,
             "trapezoid": MSO_SHAPE.TRAPEZOID, "octagon": MSO_SHAPE.OCTAGON, "plus": MSO_SHAPE.CROSS,
             "custom": MSO_SHAPE.RECTANGLE}
NUM_OUT = {QTextListFormat.ListDecimal: "arabicPeriod", QTextListFormat.ListLowerAlpha: "alphaLcPeriod",
           QTextListFormat.ListUpperAlpha: "alphaUcPeriod", QTextListFormat.ListLowerRoman: "romanLcPeriod",
           QTextListFormat.ListUpperRoman: "romanUcPeriod"}
BULLET_OUT = {QTextListFormat.ListDisc: "•", QTextListFormat.ListCircle: "o", QTextListFormat.ListSquare: "▪"}


def E(px):
    return Emu(int(round(px * EMU)))


def _rgb(hexstr):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr.lstrip("#").upper()[:6])


def _set_bg(slide, bg):
    cSld = slide._element.cSld
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    if bg is None:
        bg = ("color", "#FFFFFF")
    if bg[0] == "image":
        _, rid = slide.part.get_or_add_image_part(io.BytesIO(bg[1]))
        fill = (f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
                f'<a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    else:
        fill = f'<a:solidFill><a:srgbClr val="{bg[1].lstrip("#").upper()}"/></a:solidFill>'
    el = parse_xml(f'<p:bg {nsdecls("p", "a", "r")}><p:bgPr>{fill}<a:effectLst/></p:bgPr></p:bg>')
    cSld.insert(0, el)


def write_text(tf, doc, anchor="t", insets=None, wrap=True):
    from pptx.enum.text import MSO_ANCHOR
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}.get(anchor, MSO_ANCHOR.TOP)
    if insets:
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (E(v) for v in insets)
    block = doc.begin()
    first = True
    while block.isValid():
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _write_para(p, block)
        block = block.next()


def _write_para(p, block):
    bf = block.blockFormat()
    a = bf.alignment() & Qt.AlignHorizontal_Mask
    p.alignment = (PP_ALIGN.CENTER if a & Qt.AlignHCenter else PP_ALIGN.RIGHT if a & Qt.AlignRight
                   else PP_ALIGN.JUSTIFY if a & Qt.AlignJustify else PP_ALIGN.LEFT)
    if bf.topMargin():
        p.space_before = Pt(bf.topMargin() / PX_PER_PT)
    if bf.bottomMargin():
        p.space_after = Pt(bf.bottomMargin() / PX_PER_PT)
    if bf.lineHeightType() == QTextBlockFormat.LineHeightTypes.ProportionalHeight.value and bf.lineHeight() not in (0, 100):
        p.line_spacing = bf.lineHeight() / 90
    pPr = p._p.get_or_add_pPr()
    lst = block.textList()
    if lst is not None:
        level = max(lst.format().indent() - 1, 0)
        p.level = min(level, 8)
        marL = bf.leftMargin() + (level + 1) * INDENT_PX
        pPr.set("marL", str(int(marL * EMU)))
        pPr.set("indent", str(int(-INDENT_PX * EMU)))
        st = lst.format().style()
        if st in NUM_OUT:
            bu = parse_xml(f'<a:buAutoNum {nsdecls("a")} type="{NUM_OUT[st]}"/>')
            pPr.insert_element_before(bu, "a:tabLst", "a:defRPr", "a:extLst")
        else:
            font = parse_xml(f'<a:buFont {nsdecls("a")} typeface="Arial"/>')
            ch = parse_xml(f'<a:buChar {nsdecls("a")} char="{BULLET_OUT.get(st, chr(8226))}"/>')
            pPr.insert_element_before(font, "a:tabLst", "a:defRPr", "a:extLst")
            pPr.insert_element_before(ch, "a:tabLst", "a:defRPr", "a:extLst")
    else:
        pPr.set("marL", str(int(bf.leftMargin() * EMU)))
        pPr.set("indent", str(int(bf.textIndent() * EMU)))
        pPr.insert_element_before(parse_xml(f'<a:buNone {nsdecls("a")}/>'), "a:tabLst", "a:defRPr", "a:extLst")
    it = block.begin()
    wrote = False
    while not it.atEnd():
        frag = it.fragment()
        if frag.isValid() and frag.text():
            parts = frag.text().replace(" ", " ").split(" ")
            for i, part in enumerate(parts):
                if i:
                    p.add_line_break()
                if part:
                    r = p.add_run()
                    r.text = part
                    _write_font(r.font, frag.charFormat())
                    wrote = True
        it += 1
    end = block.charFormat()
    if not wrote:  # keep the height of empty paragraphs
        size = end.fontPointSize()
        if size:
            p.font.size = Pt(size)


def _write_font(font, cf):
    fams = cf.fontFamilies()
    if fams:
        font.name = fams[0]
    if cf.fontPointSize():
        font.size = Pt(round(cf.fontPointSize() * 2) / 2)
    font.bold = cf.fontWeight() >= QFont.DemiBold
    font.italic = cf.fontItalic()
    font.underline = cf.fontUnderline()
    fg = cf.foreground()
    if fg.style() != Qt.NoBrush:
        font.color.rgb = _rgb(fg.color().name())
    rPr = font._rPr
    if cf.fontStrikeOut():
        rPr.set("strike", "sngStrike")
    va = cf.verticalAlignment()
    if va == QTextCharFormat.AlignSuperScript:
        rPr.set("baseline", "30000")
    elif va == QTextCharFormat.AlignSubScript:
        rPr.set("baseline", "-25000")
    if cf.fontCapitalization() == QFont.AllUppercase:
        rPr.set("cap", "all")


def _apply_line(shape_line, color, width):
    if color:
        shape_line.color.rgb = _rgb(color)
        shape_line.width = E(width)
    else:
        shape_line.fill.background()


def _write_element(slide, e):
    shapes = slide.shapes
    if e.kind == "image":
        pic = shapes.add_picture(io.BytesIO(e.image), E(e.x), E(e.y), E(e.w), E(e.h))
        if any(e.crop):
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = e.crop
        shp = pic
    elif e.kind == "line":
        x1, x2 = (e.x + e.w, e.x) if e.flip_h else (e.x, e.x + e.w)
        y1, y2 = (e.y + e.h, e.y) if e.flip_v else (e.y, e.y + e.h)
        shp = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        shp.line.color.rgb = _rgb(e.line or "#000000")
        shp.line.width = E(e.line_w)
        ln = shp.line._get_or_add_ln()
        if e.arrow_start:
            ln.append(parse_xml(f'<a:headEnd {nsdecls("a")} type="triangle"/>'))
        if e.arrow_end:
            ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>'))
        return
    elif e.kind == "table":
        _write_table(shapes, e)
        return
    else:
        plain_box = e.geom == "rect" and not e.fill and not e.line
        if plain_box:
            shp = shapes.add_textbox(E(e.x), E(e.y), E(e.w), E(e.h))
        else:
            shp = shapes.add_shape(SHAPE_OUT.get(e.geom, MSO_SHAPE.RECTANGLE), E(e.x), E(e.y), E(e.w), E(e.h))
            if e.geom == "custom" and e.cust_xml:
                spPr = shp._element.spPr
                pg = spPr.find(qn("a:prstGeom"))
                cg = etree.fromstring(e.cust_xml)
                if pg is not None:
                    pg.addprevious(cg)
                    spPr.remove(pg)
            if e.fill:
                shp.fill.solid()
                shp.fill.fore_color.rgb = _rgb(e.fill)
            else:
                shp.fill.background()
            _apply_line(shp.line, e.line, e.line_w)
            style = shp._element.find(qn("p:style"))
            if style is not None:
                shp._element.remove(style)
        if e.doc is not None and not e.is_empty():
            write_text(shp.text_frame, e.doc, e.anchor, e.insets, e.wrap)
    if e.role in ("title", "body"):
        shp.name = "Ethical Title" if e.role == "title" else "Ethical Body"
    if e.rot:
        shp.rotation = e.rot
    xfrm = shp._element.spPr.find(qn("a:xfrm")) if hasattr(shp._element, "spPr") else None
    if xfrm is not None:
        if e.flip_h:
            xfrm.set("flipH", "1")
        if e.flip_v:
            xfrm.set("flipV", "1")


def table_of(doc):
    b = doc.begin()
    while b.isValid():
        t = QTextCursor(b).currentTable()
        if t is not None:
            return t
        b = b.next()
    return None


def _write_table(shapes, e):
    t = table_of(e.doc)
    if t is None:
        return
    rows, cols = t.rows(), t.columns()
    gf = shapes.add_table(rows, cols, E(e.x), E(e.y), E(e.w), E(e.h))
    table = gf.table
    tblPr = table._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    cons = t.format().columnWidthConstraints()
    if len(cons) == cols:
        for c, con in enumerate(cons):
            table.columns[c].width = E(con.rawValue())
    for r in range(rows):
        table.rows[r].height = E(e.h / rows)
    done = set()
    for r in range(rows):
        for c in range(cols):
            cell = t.cellAt(r, c)
            if (cell.row(), cell.column()) in done:
                continue
            done.add((cell.row(), cell.column()))
            pc = table.cell(r, c)
            if cell.rowSpan() > 1 or cell.columnSpan() > 1:
                pc.merge(table.cell(r + cell.rowSpan() - 1, c + cell.columnSpan() - 1))
            bg = cell.format().background()
            if bg.style() != Qt.NoBrush:
                pc.fill.solid()
                pc.fill.fore_color.rgb = _rgb(bg.color().name())
            else:
                pc.fill.background()
            tf = pc.text_frame
            it = cell.begin()
            first = True
            while not it.atEnd():
                b = it.currentBlock()
                if it.currentFrame() is None and b.isValid():
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    _write_para(p, b)
                it += 1


def save_pptx(pres, path):
    if pres.template:
        prs = PptxPresentation(io.BytesIO(pres.template))
        sld_ids = prs.slides._sldIdLst
        for sid in list(sld_ids):
            prs.part.drop_rel(sid.rId)
            sld_ids.remove(sid)
        all_layouts = [lay for m in prs.slide_masters for lay in m.slide_layouts]
    else:
        prs = PptxPresentation()
        all_layouts = list(prs.slide_layouts)
    prs.slide_width, prs.slide_height = E(pres.width), E(pres.height)
    blank = next((lay for lay in all_layouts if lay.name.lower() == "blank"), all_layouts[-1])
    for s in pres.slides:
        if pres.template and s.layout is not None and pres.layouts and s.layout < len(all_layouts):
            layout = all_layouts[pres.layouts[s.layout].index]
        else:
            layout = blank
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        if s.bg_override or not pres.template:
            _set_bg(slide, s.bg)
        for e in s.elements:
            if e.locked:
                continue
            if e.kind == "shape" and e.is_empty() and not e.fill and not e.line:
                continue  # empty placeholder
            _write_element(slide, e)
        if s.notes.strip():
            slide.notes_slide.notes_text_frame.text = s.notes
        if s.hidden:
            slide._element.set("show", "0")
    prs.save(path)
