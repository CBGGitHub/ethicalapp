"""Read/write Excel .xlsx and .csv files for Ethical Sheets (uses openpyxl)."""
import csv
import datetime
import io
import zipfile

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from sheet_engine import (Cell, ExcelError, Sheet, Workbook, cell_name, format_value,
                          parse_cell, parse_input, split_ref, to_serial)

DEFAULT_FONT, DEFAULT_SIZE = "Calibri", 11
DEFAULT_COL_PX, DEFAULT_ROW_PX = 64, 20
# Office default theme colours (lt1, dk1, lt2, dk2, accent1..6)
THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]


def width_to_px(w):
    return int(round(w * 7 + 5))


def px_to_width(px):
    return max(0.5, (px - 5) / 7)


def _tint(hexrgb, tint):
    if not tint:
        return hexrgb
    rgb = [int(hexrgb[i:i + 2], 16) for i in (0, 2, 4)]
    if tint < 0:
        rgb = [int(v * (1 + tint)) for v in rgb]
    else:
        rgb = [int(v + (255 - v) * tint) for v in rgb]
    return "".join(f"{v:02X}" for v in rgb)


def color_hex(color):
    """openpyxl Color -> '#RRGGBB' or None."""
    if color is None:
        return None
    try:
        if color.type == "rgb" and isinstance(color.rgb, str):
            if color.rgb in ("00000000",):
                return None
            return "#" + color.rgb[-6:].upper()
        if color.type == "theme" and color.theme is not None and color.theme < len(THEME):
            return "#" + _tint(THEME[color.theme], color.tint)
        if color.type == "indexed" and color.indexed is not None and color.indexed < len(COLOR_INDEX):
            if color.indexed in (64, 65):  # system foreground/background
                return None
            return "#" + COLOR_INDEX[color.indexed][-6:]
    except Exception:
        pass
    return None


def style_from(cell):
    st = {}
    f = cell.font
    if f is not None:
        if f.b:
            st["bold"] = True
        if f.i:
            st["italic"] = True
        if f.u and f.u != "none":
            st["underline"] = True
        if f.strike:
            st["strike"] = True
        if f.name and f.name != DEFAULT_FONT:
            st["font"] = f.name
        if f.sz and float(f.sz) != DEFAULT_SIZE:
            st["size"] = float(f.sz)
        c = color_hex(f.color)
        if c and c != "#000000":
            st["color"] = c
    fill = cell.fill
    if fill is not None and fill.patternType == "solid":
        c = color_hex(fill.fgColor)
        if c:
            st["fill"] = c
    al = cell.alignment
    if al is not None:
        if al.horizontal in ("left", "center", "right", "justify", "centerContinuous"):
            st["halign"] = {"centerContinuous": "center", "justify": "left"}.get(al.horizontal, al.horizontal)
        if al.vertical in ("top", "center"):
            st["valign"] = al.vertical
        if al.wrap_text:
            st["wrap"] = True
    if cell.number_format and cell.number_format != "General":
        st["numfmt"] = cell.number_format
    b = cell.border
    if b is not None:
        sides = {n: getattr(b, n).style for n in ("left", "right", "top", "bottom")
                 if getattr(b, n) is not None and getattr(b, n).style}
        if sides:
            st["border"] = sides
    return st


def _plain(v):
    if isinstance(v, ArrayFormula):
        return v.text if v.text.startswith("=") else "=" + v.text
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return float(to_serial(v))
    if isinstance(v, datetime.timedelta):
        return v.total_seconds() / 86400
    if isinstance(v, str) and v.startswith("#") and v.upper() in ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!"):
        return ExcelError(v.upper())
    if v is not None and not isinstance(v, (int, float, str, bool)):
        return str(v)
    return v


def load_xlsx(path):
    """Returns (Workbook, list of warning strings)."""
    warnings = []
    try:
        names = zipfile.ZipFile(path).namelist()
        if any(n.startswith("xl/pivot") for n in names):
            warnings.append("pivot tables")
        if any(n.endswith(".bin") and "vba" in n.lower() for n in names):
            warnings.append("macros")
    except zipfile.BadZipFile:
        raise ValueError("This isn't a valid .xlsx file.")
    xw = openpyxl.load_workbook(path)
    try:
        vw = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        vw = None
    wb = Workbook()
    wb.sheets = []
    for ws in xw.worksheets:
        sh = Sheet(ws.title)
        vs = vw[ws.title] if vw is not None and ws.title in vw.sheetnames else None
        for row in ws.iter_rows():
            for c in row:
                if c.__class__.__name__ == "MergedCell":
                    continue
                v = _plain(c.value)
                st = style_from(c) if c.has_style else {}
                if v is None and not st and not c.comment:
                    continue
                cell = Cell(v, st)
                if isinstance(v, str) and v.startswith("=") and vs is not None:
                    cell.cached = _plain(vs[c.coordinate].value)
                if c.comment:
                    cell.comment = c.comment.text
                sh.cells[(c.row - 1, c.column - 1)] = cell
        for dim in ws.column_dimensions.values():
            if dim.width and (dim.customWidth or dim.width != 13):
                for col in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
                    if col <= 16384:
                        sh.col_widths[col - 1] = width_to_px(dim.width)
            if dim.hidden:
                for col in range(dim.min or 1, (dim.max or dim.min or 1) + 1):
                    sh.col_widths[col - 1] = 0
        for r, dim in ws.row_dimensions.items():
            if dim.height:
                sh.row_heights[r - 1] = int(round(dim.height * 96 / 72))
            if dim.hidden:
                sh.row_heights[r - 1] = 0
        sh.merges = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1)
                     for m in ws.merged_cells.ranges]
        if ws.freeze_panes:
            sh.freeze = parse_cell(ws.freeze_panes)
        charts_ok = _read_charts(ws, sh)
        images_ok = _read_images(ws, sh)
        if not charts_ok and "charts" not in warnings:
            warnings.append("charts")
        if not images_ok and "pictures" not in warnings:
            warnings.append("pictures")
        if ws.conditional_formatting and "conditional formatting" not in warnings:
            warnings.append("conditional formatting")
        if ws.data_validations and ws.data_validations.dataValidation and "dropdown lists" not in warnings:
            warnings.append("dropdown lists / data validation")
        wb.sheets.append(sh)
    try:
        for name, dn in xw.defined_names.items():
            if dn.attr_text and not name.startswith("_xlnm"):
                wb.names[name.upper()] = dn.attr_text
    except Exception:
        pass
    if not wb.sheets:
        wb.sheets = [Sheet("Sheet1")]
    wb.recalc_all()
    return wb, warnings


_style_cache = {}


def _xl_styles(st):
    key = tuple(sorted((k, str(v)) for k, v in st.items()))
    if key in _style_cache:
        return _style_cache[key]
    font = Font(name=st.get("font", DEFAULT_FONT), sz=st.get("size", DEFAULT_SIZE), b=st.get("bold", False),
                i=st.get("italic", False), u="single" if st.get("underline") else None,
                strike=st.get("strike", False),
                color=st["color"].lstrip("#") if st.get("color") else None)
    fill = PatternFill("solid", fgColor=st["fill"].lstrip("#")) if st.get("fill") else None
    align = None
    if st.get("halign") or st.get("valign") or st.get("wrap"):
        align = Alignment(horizontal=st.get("halign"), vertical=st.get("valign"), wrap_text=st.get("wrap"))
    border = None
    if st.get("border"):
        border = Border(**{s: Side(style=v) for s, v in st["border"].items()})
    out = (font, fill, align, border, st.get("numfmt"))
    _style_cache[key] = out
    return out


def save_xlsx(wb, path):
    xw = openpyxl.Workbook()
    xw.remove(xw.active)
    for sh in wb.sheets:
        ws = xw.create_sheet(sh.name[:31])
        merged_hidden = {(r, c) for r1, c1, r2, c2 in sh.merges
                         for r in range(r1, r2 + 1) for c in range(c1, c2 + 1) if (r, c) != (r1, c1)}
        for (r, c), cell in sorted(sh.cells.items()):
            if (r, c) in merged_hidden:
                continue
            v = cell.value
            if isinstance(v, ExcelError):
                v = v.code
            xc = ws.cell(row=r + 1, column=c + 1, value=v)
            if isinstance(v, str) and not v.startswith("=") and v.startswith("#"):
                xc.data_type = "s"
            if cell.style:
                font, fill, align, border, numfmt = _xl_styles(cell.style)
                xc.font = font
                if fill:
                    xc.fill = fill
                if align:
                    xc.alignment = align
                if border:
                    xc.border = border
                if numfmt:
                    xc.number_format = numfmt
            if cell.comment:
                xc.comment = Comment(cell.comment, "Ethical Sheets")
        for col, px in sh.col_widths.items():
            dim = ws.column_dimensions[get_column_letter(col + 1)]
            if px == 0:
                dim.hidden = True
            else:
                dim.width = px_to_width(px)
        for row, px in sh.row_heights.items():
            if px == 0:
                ws.row_dimensions[row + 1].hidden = True
            else:
                ws.row_dimensions[row + 1].height = px * 72 / 96
        for r1, c1, r2, c2 in sh.merges:
            ws.merge_cells(start_row=r1 + 1, start_column=c1 + 1, end_row=r2 + 1, end_column=c2 + 1)
        if sh.freeze:
            ws.freeze_panes = cell_name(*sh.freeze)
        for ch in sh.charts:
            _add_chart(ws, ch)
        for im in sh.images:
            _add_image(ws, im)
    try:
        from openpyxl.workbook.defined_name import DefinedName
        for name, ref in wb.names.items():
            xw.defined_names[name] = DefinedName(name, attr_text=ref)
    except Exception:
        pass
    xw.save(path)


def _add_chart(ws, ch):
    r1, c1, r2, c2 = ch["ref"]
    kind = ch.get("type", "column")
    if kind == "pie":
        chart = PieChart()
    elif kind == "line":
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "bar" if kind == "bar" else "col"
    chart.title = ch.get("title") or None
    data = Reference(ws, min_col=c1 + 2, min_row=r1 + 1, max_col=c2 + 1, max_row=r2 + 1)
    cats = Reference(ws, min_col=c1 + 1, min_row=r1 + 2, max_row=r2 + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = ch.get("w", 480) * 2.54 / 96
    chart.height = ch.get("h", 288) * 2.54 / 96
    ws.add_chart(chart, cell_name(ch.get("row", 0), ch.get("col", 0)))


def load_csv(path):
    wb = Workbook()
    sh = wb.sheets[0]
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        for r, row in enumerate(csv.reader(fh, dialect)):
            for c, text in enumerate(row):
                v, fmt = parse_input(text)
                if v is not None:
                    sh.cells[(r, c)] = Cell(v, {"numfmt": fmt} if fmt and fmt != "@" else {})
    wb.recalc_all()
    return wb


def save_csv(wb, sheet, path):
    rows, cols = sheet.max_row() + 1, sheet.max_col() + 1
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for r in range(rows):
            w.writerow([format_value(wb.value(sheet, r, c),
                                     (sheet.get(r, c).style.get("numfmt") if sheet.get(r, c) else None) or "General")
                        for c in range(cols)])


CHART_KINDS = {"BarChart": "column", "BarChart3D": "column", "LineChart": "line", "LineChart3D": "line",
               "PieChart": "pie", "PieChart3D": "pie", "DoughnutChart": "pie"}


def _anchor(obj):
    a = obj.anchor
    frm = getattr(a, "_from", None)
    row, col = (frm.row, frm.col) if frm is not None else (0, 0)
    w = h = None
    ext = getattr(a, "ext", None)
    if ext is not None and getattr(ext, "width", None):
        w, h = ext.width / 9525, ext.height / 9525
    return row, col, w, h


def _ref_cols(ref):
    _, parts = split_ref(ref)
    (_, c1, _, r1), (_, c2, _, r2) = parts[0], parts[-1]
    return c1, c2, r1, r2


def _chart_spec(ch):
    kind = CHART_KINDS[type(ch).__name__]
    if kind == "column" and getattr(ch, "type", "col") == "bar":
        kind = "bar"
    cols, rows, cat_col = [], None, None
    for s in ch.series:
        c1, c2, r1, r2 = _ref_cols(s.val.numRef.f)
        if c1 != c2 or r1 is None:
            raise ValueError("series is not a single column")
        cols.append(c1)
        rows = (r1, r2)
        if s.cat is not None:
            cref = (s.cat.numRef or s.cat.strRef).f
            cat_col = _ref_cols(cref)[0]
    if not cols or cat_col is None or sorted(cols) != list(range(cat_col + 1, cat_col + 1 + len(cols))):
        raise ValueError("unsupported chart layout")
    title = ""
    try:
        title = "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or []))
    except AttributeError:
        pass
    row, col, w, h = _anchor(ch)
    return {"type": kind, "ref": (rows[0] - 1, cat_col, rows[1], max(cols)), "title": title,
            "row": row, "col": col, "w": int(w or ch.width * 96 / 2.54), "h": int(h or ch.height * 96 / 2.54)}


def _read_charts(ws, sheet):
    ok = True
    for ch in getattr(ws, "_charts", []):
        try:
            sheet.charts.append(_chart_spec(ch))
        except Exception:
            ok = False
    return ok


def _read_images(ws, sheet):
    ok = True
    for img in getattr(ws, "_images", []):
        try:
            row, col, w, h = _anchor(img)
            sheet.images.append({"data": img._data(), "row": row, "col": col,
                                 "w": int(w or img.width), "h": int(h or img.height)})
        except Exception:
            ok = False
    return ok


def _add_image(ws, im):
    from openpyxl.drawing.image import Image as XLImage
    x = XLImage(io.BytesIO(im["data"]))
    x.width, x.height = im["w"], im["h"]
    ws.add_image(x, cell_name(im.get("row", 0), im.get("col", 0)))
