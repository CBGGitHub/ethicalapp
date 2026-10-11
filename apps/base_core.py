"""Ethical Database core: SQLite databases with Access-style field types, relationships, queries, reports.

A database is one SQLite file (.db). Field types are stored as SQLite *declared types* that SQLite and
other tools understand: AUTONUMBER is `INTEGER PRIMARY KEY AUTOINCREMENT`, the rest are TEXT, LONGTEXT,
INTEGER, DECIMAL, CURRENCY, DATE, DATETIME, BOOLEAN. Saved queries are SQL VIEWs; saved reports live in
the `_eb_reports` table (JSON).
"""
import csv
import datetime
import html
import json
import os
import re
import sqlite3
from dataclasses import dataclass, field

FIELD_TYPES = ["AutoNumber", "Short Text", "Long Text", "Whole Number", "Decimal Number", "Currency",
               "Date", "Date & Time", "Yes/No"]
TYPE_SQL = {"Short Text": "TEXT", "Long Text": "LONGTEXT", "Whole Number": "INTEGER", "Decimal Number": "DECIMAL",
            "Currency": "CURRENCY", "Date": "DATE", "Date & Time": "DATETIME", "Yes/No": "BOOLEAN",
            "AutoNumber": "INTEGER"}
INTERNAL_TABLES = ("_eb_reports", "sqlite_sequence")


def kind_from_decl(decl, pk=False, autoinc=False):
    d = (decl or "").upper()
    if pk and d.startswith("INT") and autoinc:
        return "AutoNumber"
    if "LONGTEXT" in d or "MEMO" in d or "CLOB" in d:
        return "Long Text"
    if "BOOL" in d or d == "BIT" or "YESNO" in d:
        return "Yes/No"
    if "DATETIME" in d or "TIMESTAMP" in d:
        return "Date & Time"
    if d.startswith("DATE"):
        return "Date"
    if "CURRENCY" in d or "MONEY" in d:
        return "Currency"
    if "INT" in d:
        return "Whole Number"
    if any(k in d for k in ("REAL", "FLOA", "DOUB", "DEC", "NUM")):
        return "Decimal Number"
    return "Short Text"


def q(name):
    """Quote an identifier."""
    return '"' + str(name).replace('"', '""') + '"'


def lit(v):
    """SQL literal for a Python value (used where parameters aren't possible, e.g. in VIEWs)."""
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"


@dataclass
class Field:
    name: str
    kind: str = "Short Text"
    required: bool = False
    unique: bool = False
    default: str = ""
    pk: bool = False
    original: str = None        # name in the existing table (for renames during redesign)


@dataclass
class Relation:
    child: str
    child_field: str
    parent: str
    parent_field: str
    cascade_delete: bool = False
    cascade_update: bool = True


# ====================================================================== database

class Database:
    def __init__(self, path):
        self.path = path
        self.con = sqlite3.connect(path, detect_types=0)
        self.con.execute("PRAGMA foreign_keys = ON")
        self.con.create_function("REGEXP", 2, lambda p, s: 1 if s is not None and re.search(p, str(s), re.I) else 0)
        self.con.execute("CREATE TABLE IF NOT EXISTS _eb_reports (name TEXT PRIMARY KEY, spec TEXT)")
        self.con.commit()

    def close(self):
        self.con.close()

    # ---- schema
    def tables(self):
        rows = self.con.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name COLLATE NOCASE").fetchall()
        return [r[0] for r in rows if r[0] not in INTERNAL_TABLES and not r[0].startswith("sqlite_")]

    def views(self):
        return [r[0] for r in self.con.execute(
            "SELECT name FROM sqlite_master WHERE type='view' ORDER BY name COLLATE NOCASE")]

    def view_sql(self, name):
        r = self.con.execute("SELECT sql FROM sqlite_master WHERE type='view' AND name=?", (name,)).fetchone()
        if not r:
            return ""
        m = re.match(r"\s*CREATE\s+VIEW\s+(?:\"(?:[^\"]|\"\")*\"|\S+)\s+AS\s+(.*)$", r[0], re.I | re.S)
        return m.group(1).strip() if m else r[0]

    def is_view(self, name):
        return name in self.views()

    def fields(self, table):
        info = self.con.execute(f"PRAGMA table_info({q(table)})").fetchall()
        create = self.con.execute("SELECT sql FROM sqlite_master WHERE name=?", (table,)).fetchone()
        sql = (create[0] or "") if create else ""
        autoinc = "AUTOINCREMENT" in sql.upper()
        uniques = set()
        for idx in self.con.execute(f"PRAGMA index_list({q(table)})").fetchall():
            if idx[2]:  # unique
                cols = self.con.execute(f"PRAGMA index_info({q(idx[1])})").fetchall()
                if len(cols) == 1:
                    uniques.add(cols[0][2])
        out = []
        for cid, name, decl, notnull, dflt, pk in info:
            kind = kind_from_decl(decl, bool(pk), autoinc and bool(pk))
            if pk and kind != "AutoNumber" and (decl or "").upper().startswith("INT") and \
                    sum(1 for r in info if r[5]) == 1:
                kind = "AutoNumber"   # INTEGER PRIMARY KEY behaves like an autonumber
            d = "" if dflt is None else str(dflt).strip("'")
            out.append(Field(name, kind, bool(notnull), name in uniques, d, bool(pk), name))
        return out

    def columns(self, source):
        cur = self.con.execute(f"SELECT * FROM {q(source)} LIMIT 0")
        return [d[0] for d in cur.description]

    def kinds(self, source):
        """Field kinds for a table or view (views: from the underlying declared types)."""
        if not self.is_view(source):
            return {f.name: f.kind for f in self.fields(source)}
        info = self.con.execute(f"PRAGMA table_info({q(source)})").fetchall()
        return {r[1]: kind_from_decl(r[2]) for r in info}

    def relations(self):
        rels = []
        for t in self.tables():
            for row in self.con.execute(f"PRAGMA foreign_key_list({q(t)})").fetchall():
                _id, _seq, parent, frm, to, on_update, on_delete, _match = row
                rels.append(Relation(t, frm, parent, to or self.primary_key(parent),
                                     on_delete == "CASCADE", on_update == "CASCADE"))
        return rels

    def primary_key(self, table):
        for f in self.fields(table):
            if f.pk:
                return f.name
        return "rowid"

    def lookup_for(self, table, column):
        """(parent table, parent key, display column) if column is a foreign key."""
        for r in self.relations():
            if r.child == table and r.child_field == column:
                display = next((f.name for f in self.fields(r.parent) if f.kind in ("Short Text", "Long Text")),
                               r.parent_field)
                return r.parent, r.parent_field, display
        return None

    # ---- create / redesign tables
    def _column_sql(self, f):
        if f.kind == "AutoNumber":
            return f"{q(f.name)} INTEGER PRIMARY KEY AUTOINCREMENT"
        s = f"{q(f.name)} {TYPE_SQL[f.kind]}"
        if f.pk:
            s += " PRIMARY KEY"
        if f.required:
            s += " NOT NULL"
        if f.unique and not f.pk:
            s += " UNIQUE"
        if f.default != "":
            d = f.default
            if f.kind in ("Whole Number", "Decimal Number", "Currency"):
                try:
                    float(d)
                except ValueError:
                    raise ValueError(f"The default for {f.name} must be a number.")
            elif f.kind == "Yes/No":
                d = "1" if str(d).lower() in ("1", "yes", "true") else "0"
            elif d.upper() in ("DATE()", "NOW()", "TODAY"):
                d = "(date('now','localtime'))" if f.kind == "Date" else "(datetime('now','localtime'))"
            else:
                d = lit(d)
            s += f" DEFAULT {d}"
        return s

    def _table_sql(self, name, fields, relations):
        cols = [self._column_sql(f) for f in fields]
        for r in relations:
            cols.append(f"FOREIGN KEY ({q(r.child_field)}) REFERENCES {q(r.parent)} ({q(r.parent_field)})"
                        + (" ON DELETE CASCADE" if r.cascade_delete else " ON DELETE RESTRICT")
                        + (" ON UPDATE CASCADE" if r.cascade_update else ""))
        return f"CREATE TABLE {q(name)} (\n  " + ",\n  ".join(cols) + "\n)"

    def create_table(self, name, fields):
        if not name.strip():
            raise ValueError("Give the table a name.")
        if name in self.tables() or name in self.views():
            raise ValueError(f"There's already a table or query called {name}.")
        self._check_fields(fields)
        self.con.execute(self._table_sql(name, fields, []))
        self.con.commit()

    @staticmethod
    def _check_fields(fields):
        names = [f.name.strip() for f in fields]
        if not fields or any(not n for n in names):
            raise ValueError("Every field needs a name.")
        if len({n.lower() for n in names}) != len(names):
            raise ValueError("Two fields have the same name.")
        if sum(1 for f in fields if f.kind == "AutoNumber") > 1:
            raise ValueError("A table can only have one AutoNumber field.")

    def redesign(self, table, fields, new_name=None, relations=None):
        """Rebuild a table with new fields/relations, keeping the data (the SQLite 12-step method)."""
        self._check_fields(fields)
        new_name = new_name or table
        if relations is None:
            relations = [r for r in self.relations() if r.child == table]
            keep = {f.original for f in fields if f.original}
            relations = [r for r in relations if r.child_field in keep]
            ren = {f.original: f.name for f in fields if f.original}
            for r in relations:
                r.child_field = ren.get(r.child_field, r.child_field)
                r.child = new_name
        old_cols = {f.name for f in self.fields(table)}
        mapping = [(f.name, f.original) for f in fields if f.original in old_cols]
        indexes = [r[0] for r in self.con.execute(
            "SELECT sql FROM sqlite_master WHERE type='index' AND tbl_name=? AND sql IS NOT NULL", (table,))]
        tmp = f"_eb_new_{new_name}"
        con = self.con
        con.commit()
        con.execute("PRAGMA foreign_keys = OFF")
        con.execute("PRAGMA legacy_alter_table = ON")
        try:
            con.execute("BEGIN")
            con.execute(self._table_sql(tmp, fields, relations))
            if mapping:
                dst = ", ".join(q(n) for n, _ in mapping)
                src = ", ".join(q(o) for _, o in mapping)
                con.execute(f"INSERT INTO {q(tmp)} ({dst}) SELECT {src} FROM {q(table)}")
            con.execute(f"DROP TABLE {q(table)}")
            con.execute(f"ALTER TABLE {q(tmp)} RENAME TO {q(new_name)}")
            for sql in indexes:
                try:
                    con.execute(sql.replace(q(table), q(new_name)))
                except sqlite3.Error:
                    pass
            bad = con.execute("PRAGMA foreign_key_check").fetchall()
            if bad:
                raise ValueError(f"{len(bad)} record(s) break a relationship (e.g. an order pointing to a "
                                 "customer that doesn't exist). Fix those records first.")
            con.execute("COMMIT")
        except Exception:
            con.execute("ROLLBACK")
            raise
        finally:
            con.execute("PRAGMA legacy_alter_table = OFF")
            con.execute("PRAGMA foreign_keys = ON")
        renames = {f.original: f.name for f in fields if f.original and f.original != f.name}
        if renames or new_name != table:
            self._fix_views(table, new_name, renames)

    def _fix_views(self, table, new_table, renames):
        """Saved queries refer to tables/fields by name: follow renames so they keep working."""
        for v in self.views():
            sql = self.view_sql(v)
            if q(table) not in sql and table not in sql:
                continue
            new = sql
            for old, nn in renames.items():
                new = new.replace(f"{q(table)}.{q(old)}", f"{q(new_table)}.{q(nn)}").replace(q(old), q(nn))
            new = new.replace(q(table), q(new_table))
            if new != sql:
                try:
                    self.con.execute(f"DROP VIEW {q(v)}")
                    self.con.execute(f"CREATE VIEW {q(v)} AS {new}")
                    self.con.commit()
                except sqlite3.Error:
                    self.con.rollback()

    def broken_views(self):
        bad = []
        for v in self.views():
            try:
                self.con.execute(f"SELECT * FROM {q(v)} LIMIT 0")
            except sqlite3.Error:
                bad.append(v)
        return bad

    def rename_object(self, old, new):
        if self.is_view(old):
            sql = self.view_sql(old)
            self.con.execute(f"DROP VIEW {q(old)}")
            self.con.execute(f"CREATE VIEW {q(new)} AS {sql}")
        else:
            self.con.execute(f"ALTER TABLE {q(old)} RENAME TO {q(new)}")
        self.con.commit()

    def drop(self, name):
        if self.is_view(name):
            self.con.execute(f"DROP VIEW {q(name)}")
        else:
            self.con.execute(f"DROP TABLE {q(name)}")
        self.con.commit()

    def add_relation(self, rel):
        fields = self.fields(rel.child)
        rels = [r for r in self.relations() if r.child == rel.child and r.child_field != rel.child_field] + [rel]
        self.redesign(rel.child, fields, relations=rels)

    def remove_relation(self, rel):
        rels = [r for r in self.relations() if r.child == rel.child and r.child_field != rel.child_field]
        self.redesign(rel.child, self.fields(rel.child), relations=rels)

    # ---- data
    def count(self, source, where="", params=()):
        return self.con.execute(f"SELECT COUNT(*) FROM {q(source)} {where}", params).fetchone()[0]

    def save_query(self, name, sql):
        if name in self.tables():
            raise ValueError("There's already a table with that name.")
        sql = sql.strip().rstrip(";")
        self.con.execute(f"SELECT * FROM ({sql}) LIMIT 0")   # validate first
        if name in self.views():
            self.con.execute(f"DROP VIEW {q(name)}")
        self.con.execute(f"CREATE VIEW {q(name)} AS {sql}")
        self.con.commit()

    def compact(self):
        self.con.commit()
        ok = self.con.execute("PRAGMA integrity_check").fetchone()[0]
        self.con.execute("VACUUM")
        return ok

    # ---- reports
    def reports(self):
        return [r[0] for r in self.con.execute("SELECT name FROM _eb_reports ORDER BY name COLLATE NOCASE")]

    def report(self, name):
        r = self.con.execute("SELECT spec FROM _eb_reports WHERE name=?", (name,)).fetchone()
        return json.loads(r[0]) if r else None

    def save_report(self, name, spec):
        self.con.execute("INSERT OR REPLACE INTO _eb_reports (name, spec) VALUES (?, ?)", (name, json.dumps(spec)))
        self.con.commit()

    def delete_report(self, name):
        self.con.execute("DELETE FROM _eb_reports WHERE name=?", (name,))
        self.con.commit()


# ====================================================================== values

def to_db(value, kind):
    """Text typed by the user -> value to store (raises ValueError with a friendly message)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    s = str(value).strip()
    if s == "":
        return None
    if kind in ("Whole Number", "AutoNumber"):
        try:
            return int(float(s.replace(",", "")))
        except ValueError:
            raise ValueError("Please type a whole number.")
    if kind in ("Decimal Number", "Currency"):
        try:
            return round(float(s.replace(",", "").replace("$", "").replace("€", "").replace("£", "")), 4 if kind == "Currency" else 10)
        except ValueError:
            raise ValueError("Please type a number.")
    if kind == "Yes/No":
        return 1 if s.lower() in ("1", "yes", "true", "y", "x", "✓") else 0
    if kind in ("Date", "Date & Time"):
        d = parse_date(s)
        if d is None:
            raise ValueError("Please type a date, e.g. 2026-10-11 or 10/11/2026.")
        return d.strftime("%Y-%m-%d") if kind == "Date" else d.strftime("%Y-%m-%d %H:%M:%S")
    return s


DATE_FORMATS = ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%m/%d/%Y", "%m/%d/%Y %H:%M", "%m/%d/%y",
                "%d.%m.%Y", "%b %d, %Y", "%B %d, %Y", "%d %b %Y", "%Y/%m/%d", "%Y-%m-%dT%H:%M:%S"]


def parse_date(s):
    if isinstance(s, (datetime.date, datetime.datetime)):
        return s if isinstance(s, datetime.datetime) else datetime.datetime(s.year, s.month, s.day)
    s = str(s).strip()
    if s.lower() in ("today", "date()"):
        return datetime.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    if s.lower() in ("now", "now()"):
        return datetime.datetime.now().replace(microsecond=0)
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def display(value, kind):
    if value is None:
        return ""
    if kind == "Currency":
        try:
            return f"${float(value):,.2f}"
        except (TypeError, ValueError):
            return str(value)
    if kind == "Decimal Number" and isinstance(value, float):
        return f"{value:,.10g}" if abs(value) < 1e15 else str(value)
    if kind == "Date":
        d = parse_date(value)
        return d.strftime("%Y-%m-%d") if d else str(value)
    if kind == "Yes/No":
        return "Yes" if value in (1, "1", True) else "No"
    if isinstance(value, bytes):
        return f"<{len(value)} bytes>"
    if isinstance(value, float):
        return f"{value:,.2f}" if kind != "Decimal Number" and round(value, 2) != value or abs(value) >= 1000             else f"{value:.10g}"
    return str(value)


# ====================================================================== query designer → SQL

def criteria_sql(col, text):
    """Access-style criteria ('>100', 'Like "A*"', 'Between 1 And 5', 'Is Null', 'Smith') -> SQL condition."""
    t = text.strip()
    if not t:
        return None

    def val(v):
        v = v.strip()
        if v.startswith("#") and v.endswith("#"):
            d = parse_date(v[1:-1])
            return lit(d.strftime("%Y-%m-%d")) if d else lit(v[1:-1])
        if len(v) >= 2 and v[0] == v[-1] and v[0] in "\"'":
            return lit(v[1:-1])
        try:
            float(v)
            return v
        except ValueError:
            pass
        if v.lower() in ("yes", "true"):
            return "1"
        if v.lower() in ("no", "false"):
            return "0"
        return lit(v)
    low = t.lower()
    if low in ("is null", "null", "is empty"):
        return f"({col} IS NULL OR {col} = '')"
    if low in ("is not null", "not null", "is not empty"):
        return f"({col} IS NOT NULL AND {col} <> '')"
    m = re.match(r"(?i)^(not\s+)?like\s+(.+)$", t)
    if m:
        pat = m.group(2).strip().strip("\"'").replace("*", "%").replace("?", "_")
        return f"{col} {'NOT ' if m.group(1) else ''}LIKE {lit(pat)}"
    m = re.match(r"(?i)^between\s+(.+?)\s+and\s+(.+)$", t)
    if m:
        return f"{col} BETWEEN {val(m.group(1))} AND {val(m.group(2))}"
    m = re.match(r"^(<=|>=|<>|!=|=|<|>)\s*(.+)$", t)
    if m:
        op = "<>" if m.group(1) == "!=" else m.group(1)
        return f"{col} {op} {val(m.group(2))}"
    if " or " in low:
        parts = re.split(r"(?i)\s+or\s+", t)
        return "(" + " OR ".join(criteria_sql(col, p) for p in parts) + ")"
    if "*" in t or "?" in t:
        return f"{col} LIKE {lit(t.strip(chr(34)).replace('*', '%').replace('?', '_'))}"
    return f"{col} = {val(t)}"


def build_query(db, base, joins, rows, distinct=False, totals=False):
    """rows: [dict(table, field, show, sort, criteria, total)] -> SQL text.
    joins: [Relation] linking base to other tables."""
    sel, where, order, group = [], [], [], []
    for r in rows:
        col = f"{q(r['table'])}.{q(r['field'])}" if r["field"] != "*" else f"{q(r['table'])}.*"
        agg = r.get("total", "Group By") if totals else None
        expr = col
        if agg and agg not in ("Group By", "Where") and r["field"] != "*":
            fn = {"Sum": "SUM", "Average": "AVG", "Count": "COUNT", "Min": "MIN", "Max": "MAX"}[agg]
            expr = f"{fn}({col})"
        if r.get("show", True) and agg != "Where":
            alias = "" if r["field"] == "*" else (f" AS {q(agg + 'Of' + r['field'])}" if expr != col else "")
            sel.append(expr + alias)
            if totals and agg == "Group By":
                group.append(col)
        c = criteria_sql(expr if (totals and expr != col) else col, r.get("criteria", ""))
        if c:
            where.append((c, bool(totals and expr != col)))
        if r.get("sort") in ("Ascending", "Descending"):
            order.append(expr + (" DESC" if r["sort"] == "Descending" else ""))
    if not sel:
        sel = [f"{q(base)}.*"]
    sql = "SELECT " + ("DISTINCT " if distinct else "") + ", ".join(sel) + f"\nFROM {q(base)}"
    joined = {base}
    for j in joins:
        cond = f"{q(j.parent)}.{q(j.parent_field)} = {q(j.child)}.{q(j.child_field)}"
        if j.child in joined and j.parent not in joined:
            sql += f"\nLEFT JOIN {q(j.parent)} ON {cond}"
            joined.add(j.parent)
        elif j.parent in joined and j.child not in joined:
            sql += f"\nLEFT JOIN {q(j.child)} ON {cond}"
            joined.add(j.child)
    plain = [c for c, is_agg in where if not is_agg]
    having = [c for c, is_agg in where if is_agg]
    if plain:
        sql += "\nWHERE " + " AND ".join(plain)
    if group:
        sql += "\nGROUP BY " + ", ".join(group)
    if having:
        sql += "\nHAVING " + " AND ".join(having)
    if order:
        sql += "\nORDER BY " + ", ".join(order)
    return sql


# ====================================================================== import / export

def infer_kind(values):
    vals = [v for v in values if v not in (None, "")]
    if not vals:
        return "Short Text"

    def all_match(fn):
        try:
            return all(fn(v) for v in vals)
        except Exception:
            return False
    if all(isinstance(v, bool) for v in vals) or all(str(v).lower() in ("yes", "no", "true", "false") for v in vals):
        return "Yes/No"
    if all(isinstance(v, (datetime.date, datetime.datetime)) for v in vals):
        return "Date" if all(getattr(v, "hour", 0) == 0 and getattr(v, "minute", 0) == 0 for v in vals) else "Date & Time"
    if all_match(lambda v: float(str(v).replace(",", "")) == int(float(str(v).replace(",", ""))) and "." not in str(v)):
        return "Whole Number"
    if all_match(lambda v: str(v).strip().startswith(("$", "€", "£")) and float(str(v).strip()[1:].replace(",", "")) is not None):
        return "Currency"
    if all_match(lambda v: float(str(v).replace(",", "")) is not None):
        return "Decimal Number"
    if all(not isinstance(v, (int, float)) and parse_date(v) is not None for v in vals):
        return "Date"
    if any(len(str(v)) > 255 or "\n" in str(v) for v in vals):
        return "Long Text"
    return "Short Text"


def clean_name(name, used):
    n = re.sub(r"\s+", " ", str(name or "").strip()) or "Field"
    base, i = n, 2
    while n.lower() in used:
        n = f"{base} {i}"
        i += 1
    used.add(n.lower())
    return n


def import_rows(db, table, header, rows, add_id=True):
    """Create a new table from rows (list of lists) and insert them. Returns record count."""
    used = set()
    names = [clean_name(h, used) for h in header]
    kinds = [infer_kind([r[i] if i < len(r) else None for r in rows[:2000]]) for i in range(len(names))]
    fields = []
    if add_id and "id" not in used:
        fields.append(Field("ID", "AutoNumber", pk=True))
    fields += [Field(n, k) for n, k in zip(names, kinds)]
    db.create_table(table, fields)
    cols = ", ".join(q(n) for n in names)
    marks = ", ".join("?" for _ in names)
    out = []
    for r in rows:
        rec = []
        for i, k in enumerate(kinds):
            v = r[i] if i < len(r) else None
            try:
                rec.append(to_db(v.isoformat(" ") if isinstance(v, (datetime.date, datetime.datetime)) else v, k))
            except ValueError:
                rec.append(None if v in (None, "") else str(v))
        out.append(rec)
    db.con.executemany(f"INSERT INTO {q(table)} ({cols}) VALUES ({marks})", out)
    db.con.commit()
    return len(out)


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(fh, dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    return (rows[0], rows[1:]) if rows else ([], [])


def read_xlsx(path):
    """-> {sheet name: (header, rows)} using cached values."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v not in (None, "") for v in r)]
        if rows:
            width = max(len(r) for r in rows)
            rows = [r + [None] * (width - len(r)) for r in rows]
            out[ws.title] = ([str(h) if h is not None else f"Column {i + 1}" for i, h in enumerate(rows[0])], rows[1:])
    return out


def access_driver():
    try:
        import pyodbc
        return next((d for d in pyodbc.drivers() if "Microsoft Access Driver" in d), None)
    except Exception:
        return None


def read_access(path):
    """-> {table: (header, rows)} via Microsoft's Access ODBC driver."""
    import pyodbc
    drv = access_driver()
    if not drv:
        raise RuntimeError("To import Access files, install Microsoft's free “Access Database Engine 2016 "
                           "Redistributable” (64-bit) from microsoft.com, then try again.")
    con = pyodbc.connect(f"DRIVER={{{drv}}};DBQ={path};")
    out = {}
    cur = con.cursor()
    names = [t.table_name for t in cur.tables(tableType="TABLE")]
    for t in names:
        cur.execute(f"SELECT * FROM [{t}]")
        header = [d[0] for d in cur.description]
        out[t] = (header, [list(r) for r in cur.fetchall()])
    con.close()
    return out


def fetch(db, source, where="", order="", params=()):
    cur = db.con.execute(f"SELECT * FROM {q(source)} {where} {order}", params)
    return [d[0] for d in cur.description], cur.fetchall()


def export_csv(db, source, path):
    header, rows = fetch(db, source)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def export_xlsx(db, sources, path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for src in sources:
        header, rows = fetch(db, src)
        kinds = db.kinds(src)
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", src)[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="7B1F3A")
        for r in rows:
            out = []
            for h, v in zip(header, r):
                k = kinds.get(h)
                if k in ("Date", "Date & Time") and v:
                    v = parse_date(v) or v
                elif k == "Yes/No" and v is not None:
                    v = bool(v)
                out.append(v)
            ws.append(out)
        for i, h in enumerate(header, start=1):
            k = kinds.get(h)
            letter = openpyxl.utils.get_column_letter(i)
            ws.column_dimensions[letter].width = max(10, min(40, len(str(h)) + 4))
            if k == "Currency":
                for c in ws[letter][1:]:
                    c.number_format = "$#,##0.00"
            elif k == "Date":
                for c in ws[letter][1:]:
                    c.number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
    wb.save(path)


# ====================================================================== reports

def report_html(db, spec):
    """spec: source, title, fields, group_by, sort, totals {field: 'Sum'|'Average'|'Count'}, landscape."""
    src = spec["source"]
    kinds = db.kinds(src)
    fields = spec.get("fields") or db.columns(src)
    group = spec.get("group_by") or ""
    order = []
    if group:
        order.append(q(group))
    if spec.get("sort"):
        order.append(q(spec["sort"]) + (" DESC" if spec.get("sort_desc") else ""))
    header, rows = fetch(db, src, order=("ORDER BY " + ", ".join(order)) if order else "")
    idx = {h: i for i, h in enumerate(header)}
    totals = spec.get("totals", {})

    def total_row(recs, label):
        cells = []
        for f in fields:
            how = totals.get(f)
            if not how:
                cells.append("")
                continue
            vals = [r[idx[f]] for r in recs if r[idx[f]] not in (None, "")]
            if how == "Count":
                v = str(len(vals))
            else:
                nums = []
                for x in vals:
                    try:
                        nums.append(float(x))
                    except (TypeError, ValueError):
                        pass
                n = sum(nums) if how == "Sum" else (sum(nums) / len(nums) if nums else 0)
                v = display(n, "Currency") if kinds.get(f) == "Currency" else f"{n:,.2f}".rstrip("0").rstrip(".")
            cells.append(f"<b>{html.escape(how)}: {html.escape(v)}</b>")
        return '<tr class="tot">' + "".join(
            f'<td align="{_align(kinds.get(f))}">{c}</td>' for c, f in zip(cells, fields)) + "</tr>"

    title = html.escape(spec.get("title") or src)
    out = [f"""<html><head><style>
      body {{ font-family: 'Segoe UI', Arial; font-size: 10pt; color: #222; }}
      h1 {{ font-size: 20pt; color: #7B1F3A; margin-bottom: 0; }}
      .sub {{ color: #777; font-size: 9pt; }}
      table {{ border-collapse: collapse; width: 100%; }}
      th {{ background: #7B1F3A; color: white; text-align: left; padding: 4px 6px; font-weight: 600; }}
      td {{ padding: 3px 6px; border-bottom: 1px solid #ddd; }}
      tr.grp td {{ background: #f3e6ea; font-weight: bold; font-size: 11pt; padding-top: 8px; }}
      tr.tot td {{ border-top: 2px solid #7B1F3A; background: #fafafa; }}
    </style></head><body>
    <h1>{title}</h1><p class="sub">{len(rows)} record(s) · {datetime.datetime.now():%B %d, %Y %H:%M}</p>
    <table cellspacing="0" cellpadding="4" width="100%"><thead><tr>""" +
           "".join(f'<th align="{_align(kinds.get(f))}">{html.escape(f)}</th>' for f in fields) + "</tr></thead><tbody>"]
    groups = []
    if group and group in idx:
        cur, bucket = object(), []
        for r in rows:
            if r[idx[group]] != cur:
                if bucket:
                    groups.append((cur, bucket))
                cur, bucket = r[idx[group]], []
            bucket.append(r)
        if bucket:
            groups.append((cur, bucket))
    else:
        groups = [(None, rows)]
    for key, recs in groups:
        if group:
            out.append(f'<tr class="grp"><td colspan="{len(fields)}">{html.escape(group)}: '
                       f'{html.escape(display(key, kinds.get(group)) or "(blank)")}  ({len(recs)})</td></tr>')
        for r in recs:
            out.append("<tr>" + "".join(
                f'<td align="{_align(kinds.get(f))}">{html.escape(display(r[idx[f]], kinds.get(f)))}</td>' for f in fields) + "</tr>")
        if group and totals:
            out.append(total_row(recs, "Subtotal"))
    if totals:
        out.append(total_row(rows, "Total"))
    out.append("</tbody></table></body></html>")
    return "".join(out)


def _align(kind):
    return "right" if kind in ("Whole Number", "Decimal Number", "Currency", "AutoNumber") else \
        "center" if kind == "Yes/No" else "left"


# ====================================================================== templates

TEMPLATES = {
    "Blank database": [],
    "Contacts": [
        ("Contacts", [Field("ID", "AutoNumber", pk=True), Field("First Name", required=True), Field("Last Name"),
                      Field("Company"), Field("Email", unique=False), Field("Phone"), Field("Address"),
                      Field("City"), Field("Birthday", "Date"), Field("Favorite", "Yes/No", default="0"),
                      Field("Notes", "Long Text")],
         [("Ada", "Lovelace", "Analytical Engines", "ada@example.com", "555-0101", "12 Byron St", "London", "1815-12-10", 1, "First programmer"),
          ("Alan", "Turing", "Bletchley Park", "alan@example.com", "555-0102", "1 Hut 8", "Bletchley", "1912-06-23", 0, ""),
          ("Grace", "Hopper", "US Navy", "grace@example.com", "555-0103", "7 Harbor Rd", "New York", "1906-12-09", 1, "Found the first bug")]),
    ],
    "Tasks": [
        ("Tasks", [Field("ID", "AutoNumber", pk=True), Field("Title", required=True), Field("Assigned To"),
                   Field("Due Date", "Date"), Field("Priority", default="Normal"), Field("Done", "Yes/No", default="0"),
                   Field("Notes", "Long Text")],
         [("Plan the launch", "Sam", "2026-10-20", "High", 0, ""), ("Order supplies", "Lee", "2026-10-14", "Normal", 1, ""),
          ("Write newsletter", "Sam", "2026-10-30", "Low", 0, "Include photos")]),
    ],
    "Inventory": [
        ("Categories", [Field("ID", "AutoNumber", pk=True), Field("Name", required=True, unique=True)],
         [("Tools",), ("Paint",), ("Garden",)]),
        ("Products", [Field("ID", "AutoNumber", pk=True), Field("Name", required=True), Field("Category ID", "Whole Number"),
                      Field("SKU", unique=True), Field("Price", "Currency"), Field("In Stock", "Whole Number", default="0"),
                      Field("Reorder Level", "Whole Number", default="5"), Field("Discontinued", "Yes/No", default="0")],
         [("Hammer", 1, "T-100", 14.99, 25, 5, 0), ("Screwdriver set", 1, "T-110", 22.5, 3, 5, 0),
          ("White paint 1L", 2, "P-200", 9.75, 40, 10, 0), ("Garden hose", 3, "G-300", 29.0, 2, 4, 0)]),
    ],
    "Customers & Orders": [
        ("Customers", [Field("ID", "AutoNumber", pk=True), Field("Name", required=True), Field("Email"), Field("Phone"),
                       Field("City")],
         [("Northwind Traders", "orders@northwind.example", "555-0200", "Seattle"),
          ("Contoso Café", "hello@contoso.example", "555-0201", "Portland")]),
        ("Products", [Field("ID", "AutoNumber", pk=True), Field("Name", required=True), Field("Price", "Currency")],
         [("Coffee beans 1kg", 18.0), ("Green tea 500g", 12.5), ("Mug", 7.0)]),
        ("Orders", [Field("ID", "AutoNumber", pk=True), Field("Customer ID", "Whole Number", required=True),
                    Field("Order Date", "Date", default="DATE()"), Field("Status", default="New")],
         [(1, "2026-10-01", "Shipped"), (2, "2026-10-08", "New")]),
        ("Order Items", [Field("ID", "AutoNumber", pk=True), Field("Order ID", "Whole Number", required=True),
                         Field("Product ID", "Whole Number", required=True), Field("Quantity", "Whole Number", default="1"),
                         Field("Unit Price", "Currency")],
         [(1, 1, 3, 18.0), (1, 3, 2, 7.0), (2, 2, 4, 12.5)]),
    ],
}
TEMPLATE_RELATIONS = {
    "Inventory": [Relation("Products", "Category ID", "Categories", "ID")],
    "Customers & Orders": [Relation("Orders", "Customer ID", "Customers", "ID"),
                           Relation("Order Items", "Order ID", "Orders", "ID", cascade_delete=True),
                           Relation("Order Items", "Product ID", "Products", "ID")],
}
TEMPLATE_QUERIES = {
    "Inventory": {"Low Stock": 'SELECT "Products"."Name", "Products"."In Stock", "Products"."Reorder Level"\n'
                               'FROM "Products"\nWHERE "Products"."In Stock" <= "Products"."Reorder Level" '
                               'AND "Products"."Discontinued" = 0\nORDER BY "Products"."In Stock"'},
    "Customers & Orders": {"Order Totals": 'SELECT "Orders"."ID" AS "Order", "Customers"."Name" AS "Customer", '
                                           '"Orders"."Order Date", SUM("Order Items"."Quantity" * "Order Items"."Unit Price") AS "Total"\n'
                                           'FROM "Orders"\nLEFT JOIN "Customers" ON "Customers"."ID" = "Orders"."Customer ID"\n'
                                           'LEFT JOIN "Order Items" ON "Order Items"."Order ID" = "Orders"."ID"\n'
                                           'GROUP BY "Orders"."ID"\nORDER BY "Orders"."Order Date" DESC'},
    "Tasks": {"Open Tasks": 'SELECT * FROM "Tasks" WHERE "Done" = 0 ORDER BY "Due Date"'},
}


def create_from_template(path, name):
    if os.path.exists(path):
        os.remove(path)
    db = Database(path)
    for table, fields, rows in TEMPLATES[name]:
        rels = [r for r in TEMPLATE_RELATIONS.get(name, []) if r.child == table]
        db.con.execute(db._table_sql(table, fields, rels))
        cols = [f for f in fields if f.kind != "AutoNumber"]
        db.con.executemany(f"INSERT INTO {q(table)} ({', '.join(q(f.name) for f in cols)}) VALUES "
                           f"({', '.join('?' for _ in cols)})", rows)
    for qname, sql in TEMPLATE_QUERIES.get(name, {}).items():
        db.con.execute(f"CREATE VIEW {q(qname)} AS {sql}")
    if name == "Customers & Orders":
        db.save_report("Orders by Customer", {"source": "Order Totals", "title": "Orders by Customer",
                                              "fields": ["Order", "Order Date", "Total"], "group_by": "Customer",
                                              "sort": "Order Date", "totals": {"Total": "Sum"}})
    db.con.commit()
    return db
