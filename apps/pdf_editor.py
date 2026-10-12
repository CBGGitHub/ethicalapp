"""PDF Editor - view, annotate, edit, fill, sign and organise PDF files."""
import io
import os
import re
import sys

import pymupdf
from PySide6.QtCore import Qt, QPointF, QRectF, QSize, QTimer
from PySide6.QtGui import (QAction, QActionGroup, QColor, QFont, QGuiApplication, QIcon, QImage,
                           QKeySequence, QPainter, QPainterPath, QPen, QPixmap)
from PySide6.QtPrintSupport import QPrintDialog, QPrinter
from PySide6.QtWidgets import (QCheckBox, QComboBox, QDialog, QDialogButtonBox, QDockWidget,
                               QFileDialog, QFormLayout, QHBoxLayout, QInputDialog, QLabel,
                               QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMenu,
                               QMessageBox, QPlainTextEdit, QPushButton, QScrollArea, QSpinBox,
                               QTabWidget, QToolBar, QVBoxLayout, QWidget)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from suite_common import (SUITE_NAME, ColorButton, font_families, font_file,  # noqa: E402
                          make_app, match_family)
from suite_common import add_guide_action  # noqa: E402
from i18n import tr  # noqa: E402

APP_NAME = "Ethical PDF Editor"

# display name -> (base14 code for insert_text regular, bold, FreeText annotation font)
FONTS = {
    "Helvetica": ("helv", "hebo", "Helv"),
    "Times": ("tiro", "tibo", "TiRo"),
    "Courier": ("cour", "cobo", "Cour"),
}
FREETEXT_FONT = {v[2]: k for k, v in FONTS.items()}

TOOLS = [  # (key, label, tooltip)
    ("select", tr("Select"), tr("Select / move annotations, select text, fill form fields")),
    ("edittext", tr("Edit Text"), tr("Click a line of existing text to change it")),
    ("addtext", tr("Add Text"), tr("Click to place a text box")),
    ("highlight", tr("Highlight"), tr("Drag over text to highlight it")),
    ("underline", tr("Underline"), tr("Drag over text to underline it")),
    ("strike", tr("Strikeout"), tr("Drag over text to strike it out")),
    ("note", tr("Note"), tr("Click to add a sticky note")),
    ("ink", tr("Draw"), tr("Freehand drawing")),
    ("rect", tr("Rectangle"), tr("Drag to draw a rectangle")),
    ("ellipse", tr("Ellipse"), tr("Drag to draw an ellipse")),
    ("line", tr("Line"), tr("Drag to draw a line")),
    ("arrow", tr("Arrow"), tr("Drag to draw an arrow")),
    ("whiteout", tr("Whiteout"), tr("Drag to paint over content with white")),
    ("redact", tr("Redact"), tr("Drag to permanently remove content (black box)")),
    ("image", tr("Image"), tr("Click or drag to insert an image")),
    ("signature", tr("Sign"), tr("Click to place your signature")),
]
TEXT_TOOLS = ("highlight", "underline", "strike")
DRAG_TOOLS = ("rect", "ellipse", "line", "arrow", "whiteout", "redact")


def rgb(c):
    return (c.redF(), c.greenF(), c.blueF())


def int_to_qcolor(v):
    return QColor((v >> 16) & 255, (v >> 8) & 255, v & 255)


def pix_to_qimage(pix):
    fmt = QImage.Format_RGBA8888 if pix.alpha else QImage.Format_RGB888
    return QImage(pix.samples, pix.width, pix.height, pix.stride, fmt).copy()


# ---------------------------------------------------------------- text selection helpers

def nearest_word(words, p):
    best, bi = None, -1
    for i, w in enumerate(words):
        dx = max(w[0] - p.x, 0, p.x - w[2])
        dy = max(w[1] - p.y, 0, p.y - w[3])
        d = dx * dx + 4 * dy * dy
        if best is None or d < best:
            best, bi = d, i
    return bi


def words_between(words, p1, p2):
    if not words:
        return []
    i, j = nearest_word(words, p1), nearest_word(words, p2)
    if i > j:
        i, j = j, i
    return words[i:j + 1]


def line_rects(sel):
    out = {}
    for w in sel:
        key, r = (w[5], w[6]), pymupdf.Rect(w[:4])
        out[key] = out[key] | r if key in out else r
    return list(out.values())


def words_text(sel):
    lines, last = [], None
    for w in sel:
        key = (w[5], w[6])
        if key != last:
            lines.append([])
            last = key
        lines[-1].append(w[4])
    return "\n".join(" ".join(line) for line in lines)


# ---------------------------------------------------------------- dialogs

class TextDialog(QDialog):
    """Text + font picker. styles=False hides bold/italic and limits fonts to the PDF built-ins
    (used for FreeText annotations, which only support those)."""

    def __init__(self, parent, title, text="", size=12, family="Arial", color=QColor("black"),
                 bold=False, italic=False, styles=True):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.edit = QPlainTextEdit(text)
        self.edit.setMinimumSize(380, 120)
        self.family = QComboBox()
        self.family.addItems(font_families() if styles else list(FONTS))
        self.family.setCurrentText(family if self.family.findText(family) >= 0
                                   else self.family.itemText(0))
        self.family.setMaxVisibleItems(20)
        self.family.currentTextChanged.connect(self._preview_font)
        self.size = QSpinBox()
        self.size.setRange(4, 144)
        self.size.setValue(int(round(size)))
        self.color = ColorButton(color, tr("Text colour"))
        self.bold = QCheckBox(tr("Bold"))
        self.bold.setChecked(bold)
        self.italic = QCheckBox(tr("Italic"))
        self.italic.setChecked(italic)
        for cb in (self.bold, self.italic):
            cb.toggled.connect(self._preview_font)
        form = QFormLayout()
        form.addRow(tr("Font:"), self.family)
        form.addRow(tr("Size:"), self.size)
        if styles:
            row = QHBoxLayout()
            row.addWidget(self.bold)
            row.addWidget(self.italic)
            row.addStretch()
            form.addRow(tr("Style:"), row)
        else:
            self.bold.setChecked(False)
            self.italic.setChecked(False)
        form.addRow(tr("Colour:"), self.color)
        self._preview_font()
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        lay = QVBoxLayout(self)
        lay.addWidget(self.edit)
        lay.addLayout(form)
        lay.addWidget(buttons)

    def _preview_font(self, *_):
        f = QFont(self.family.currentText(), 12)
        f.setBold(self.bold.isChecked())
        f.setItalic(self.italic.isChecked())
        self.edit.setFont(f)

    def values(self):
        return (self.edit.toPlainText(), self.size.value(), self.family.currentText(),
                self.color.color(), self.bold.isChecked(), self.italic.isChecked())


class SignatureCanvas(QWidget):
    def __init__(self):
        super().__init__()
        self.setFixedSize(460, 180)
        self.strokes = []

    def clear(self):
        self.strokes = []
        self.update()

    def mousePressEvent(self, e):
        self.strokes.append([e.position()])
        self.update()

    def mouseMoveEvent(self, e):
        if self.strokes:
            self.strokes[-1].append(e.position())
            self.update()

    def paintEvent(self, e):
        p = QPainter(self)
        p.setRenderHint(QPainter.Antialiasing)
        p.fillRect(self.rect(), Qt.white)
        p.setPen(QPen(QColor("#bbb"), 1, Qt.DashLine))
        p.drawLine(20, 140, self.width() - 20, 140)
        p.setPen(QPen(QColor("#102a6b"), 2.5, Qt.SolidLine, Qt.RoundCap, Qt.RoundJoin))
        for s in self.strokes:
            path = QPainterPath(s[0])
            for pt in s[1:]:
                path.lineTo(pt)
            p.drawPath(path)


class SignatureDialog(QDialog):
    def __init__(self, parent):
        super().__init__(parent)
        self.setWindowTitle(tr("Create Signature"))
        self.result_sig = None
        self.tabs = QTabWidget()
        draw = QWidget()
        self.canvas = SignatureCanvas()
        clear = QPushButton(tr("Clear"))
        clear.clicked.connect(self.canvas.clear)
        dl = QVBoxLayout(draw)
        dl.addWidget(QLabel(tr("Draw your signature with the mouse:")))
        dl.addWidget(self.canvas)
        dl.addWidget(clear, alignment=Qt.AlignLeft)
        self.tabs.addTab(draw, tr("Draw"))
        img = QWidget()
        il = QVBoxLayout(img)
        self.img_path = QLineEdit()
        self.img_path.setReadOnly(True)
        browse = QPushButton(tr("Choose image..."))
        browse.clicked.connect(self._browse)
        il.addWidget(QLabel(tr("Use a picture of your signature (PNG with transparency works best):")))
        il.addWidget(self.img_path)
        il.addWidget(browse, alignment=Qt.AlignLeft)
        il.addStretch()
        self.tabs.addTab(img, tr("Image"))
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self._ok)
        buttons.rejected.connect(self.reject)
        lay = QVBoxLayout(self)
        lay.addWidget(self.tabs)
        lay.addWidget(buttons)

    def _browse(self):
        f, _ = QFileDialog.getOpenFileName(self, tr("Signature image"), "",
                                           tr("Images (*.png *.jpg *.jpeg *.bmp)"))
        if f:
            self.img_path.setText(f)

    def _ok(self):
        if self.tabs.currentIndex() == 0:
            strokes = [[(p.x(), p.y()) for p in s] for s in self.canvas.strokes if len(s) > 1]
            if not strokes:
                return
            self.result_sig = ("ink", strokes)
        else:
            if not self.img_path.text():
                return
            self.result_sig = ("image", self.img_path.text())
        self.accept()


# ---------------------------------------------------------------- page view

class PageWidget(QWidget):
    def __init__(self, win, idx):
        super().__init__()
        self.win, self.idx = win, idx
        self.resize_to_page()

    def resize_to_page(self):
        r = self.win.doc[self.idx].rect
        z = self.win.zoom
        self.setFixedSize(max(1, int(r.width * z)), max(1, int(r.height * z)))

    def paintEvent(self, e):
        p = QPainter(self)
        p.fillRect(self.rect(), Qt.white)
        p.drawImage(QRectF(0, 0, self.width(), self.height()), self.win.page_image(self.idx))
        p.setRenderHint(QPainter.Antialiasing)
        self.win.paint_overlay(self.idx, p)

    def mousePressEvent(self, e):
        self.win.on_press(self, e)

    def mouseMoveEvent(self, e):
        self.win.on_move(self, e)

    def mouseReleaseEvent(self, e):
        self.win.on_release(self, e)

    def mouseDoubleClickEvent(self, e):
        self.win.on_double_click(self, e)

    def contextMenuEvent(self, e):
        self.win.on_context_menu(self, e)


class PageScroll(QScrollArea):
    def __init__(self, win):
        super().__init__()
        self.win = win

    def wheelEvent(self, e):
        if e.modifiers() & Qt.ControlModifier:
            self.win.set_zoom(self.win.zoom * (1.1 if e.angleDelta().y() > 0 else 1 / 1.1))
            e.accept()
        else:
            super().wheelEvent(e)


# ---------------------------------------------------------------- main window

class PdfEditor(QMainWindow):
    def __init__(self, path=None):
        super().__init__()
        self.doc = None
        self.path = None
        self.password = None
        self.dirty = False
        self.zoom = 1.25
        self.tool = "select"
        self.undo_stack, self.redo_stack = [], []
        self.img_cache, self.field_cache = {}, {}
        self.page_widgets = []
        self.drag = None
        self.sel_words = None       # (idx, [words])
        self.sel_annot = None       # (idx, xref, rect)
        self.search_hits, self.search_pos, self.search_query = [], -1, ""
        self.signature = None
        self.font_cache = {}
        self.show_fields = True
        self.thumb_queue = []

        self.resize(1300, 900)
        self.setAcceptDrops(True)
        self._build_ui()
        self.new_blank(confirm=False)
        if path:
            self.open_path(path)

    # ------------------------------------------------------------ UI
    def _build_ui(self):
        self.scroll = PageScroll(self)
        self.scroll.setWidgetResizable(True)
        self.scroll.setAlignment(Qt.AlignHCenter)
        self.container = QWidget()
        self.container.setStyleSheet("background: #6b6f76;")
        self.vbox = QVBoxLayout(self.container)
        self.vbox.setSpacing(14)
        self.vbox.setContentsMargins(20, 20, 20, 20)
        self.vbox.setAlignment(Qt.AlignHCenter | Qt.AlignTop)
        self.scroll.setWidget(self.container)
        self.scroll.verticalScrollBar().valueChanged.connect(self._track_current_page)
        self.setCentralWidget(self.scroll)

        self.thumbs = QListWidget()
        self.thumbs.setViewMode(QListWidget.IconMode)
        self.thumbs.setFlow(QListWidget.TopToBottom)
        self.thumbs.setWrapping(False)
        self.thumbs.setMovement(QListWidget.Static)
        self.thumbs.setIconSize(QSize(130, 170))
        self.thumbs.setGridSize(QSize(150, 200))
        self.thumbs.setSelectionMode(QListWidget.ExtendedSelection)
        self.thumbs.setStyleSheet("QListWidget { background: #d9dbe0; }")
        self.thumbs.setMinimumWidth(170)
        self.thumbs.itemClicked.connect(lambda it: self.goto_page(self.thumbs.row(it)))
        self.thumbs.setContextMenuPolicy(Qt.CustomContextMenu)
        self.thumbs.customContextMenuRequested.connect(
            lambda pos: self.pages_menu.exec(self.thumbs.mapToGlobal(pos)))
        dock = QDockWidget(tr("Pages"), self)
        dock.setObjectName("pages")
        dock.setWidget(self.thumbs)
        dock.setFeatures(QDockWidget.DockWidgetMovable | QDockWidget.DockWidgetClosable)
        self.addDockWidget(Qt.LeftDockWidgetArea, dock)
        self.thumb_dock = dock
        self.thumb_timer = QTimer(self)
        self.thumb_timer.timeout.connect(self._render_some_thumbs)

        self._build_actions()
        self.status = self.statusBar()

    def act(self, text, slot, shortcut=None, menu=None, tip=None):
        a = QAction(text, self)
        a.triggered.connect(slot)
        if shortcut:
            a.setShortcut(QKeySequence(shortcut))
        if tip:
            a.setStatusTip(tip)
        if menu is not None:
            menu.addAction(a)
        self.addAction(a)
        return a

    def _build_actions(self):
        mb = self.menuBar()
        f = mb.addMenu(tr("&File"))
        a_new = self.act(tr("&New Blank PDF"), self.new_blank, "Ctrl+N", f)
        self.act(tr("New PDF from &Images..."), self.new_from_images, None, f)
        a_open = self.act(tr("&Open..."), self.open_dialog, "Ctrl+O", f)
        self.act(tr("&Combine Files into PDF..."), self.combine_files, None, f)
        f.addSeparator()
        a_save = self.act(tr("&Save"), self.save, "Ctrl+S", f)
        self.act(tr("Save &As..."), self.save_as, "Ctrl+Shift+S", f)
        ex = f.addMenu(tr("&Export"))
        self.act(tr("Pages as &Images (PNG/JPG)..."), self.export_images, None, ex)
        self.act(tr("As &Word Document (.docx)..."), self.export_docx, None, ex)
        self.act(tr("As Plain &Text (.txt)..."), self.export_text, None, ex)
        f.addSeparator()
        a_print = self.act(tr("&Print..."), self.print_doc, "Ctrl+P", f)
        f.addSeparator()
        self.act(tr("E&xit"), self.close, "Alt+F4", f)

        e = mb.addMenu(tr("&Edit"))
        a_undo = self.act(tr("&Undo"), self.undo, "Ctrl+Z", e)
        a_redo = self.act(tr("&Redo"), self.redo, "Ctrl+Y", e)
        e.addSeparator()
        self.act(tr("&Copy Selected Text"), self.copy_text, "Ctrl+C", e)
        self.act(tr("&Delete Selected Annotation"), self.delete_selected_annot, "Delete", e)
        e.addSeparator()
        self.act(tr("&Find..."), lambda: (self.search_box.setFocus(), self.search_box.selectAll()),
                 "Ctrl+F", e)
        self.act(tr("Create &Signature..."), self.create_signature, None, e)

        v = mb.addMenu(tr("&View"))
        a_zin = self.act(tr("Zoom &In"), lambda: self.set_zoom(self.zoom * 1.2), "Ctrl+=", v)
        a_zout = self.act(tr("Zoom &Out"), lambda: self.set_zoom(self.zoom / 1.2), "Ctrl+-", v)
        self.act(tr("&Actual Size"), lambda: self.set_zoom(1.0), "Ctrl+0", v)
        a_fit = self.act(tr("Fit &Width"), self.fit_width, "Ctrl+1", v)
        v.addSeparator()
        hf = self.act(tr("Highlight Form &Fields"), self.toggle_fields, None, v)
        hf.setCheckable(True)
        hf.setChecked(True)
        v.addAction(self.thumb_dock.toggleViewAction())

        pm = mb.addMenu(tr("&Pages"))
        self.pages_menu = pm
        self.act(tr("Rotate &Right"), lambda: self.rotate_pages(90), "Ctrl+R", pm)
        self.act(tr("Rotate &Left"), lambda: self.rotate_pages(-90), "Ctrl+Shift+R", pm)
        pm.addSeparator()
        self.act(tr("Move Page &Up"), lambda: self.move_page(-1), None, pm)
        self.act(tr("Move Page &Down"), lambda: self.move_page(1), None, pm)
        pm.addSeparator()
        self.act(tr("Insert &Blank Page"), self.insert_blank, None, pm)
        self.act(tr("Insert Pages from &File..."), self.insert_from_file, None, pm)
        self.act(tr("Insert &Image as Page..."), self.insert_image_pages, None, pm)
        pm.addSeparator()
        self.act(tr("&Extract Selected Pages..."), self.extract_pages, None, pm)
        self.act(tr("&Split into Single Pages..."), self.split_pages, None, pm)
        pm.addSeparator()
        self.act(tr("De&lete Selected Pages"), self.delete_pages, None, pm)

        h = mb.addMenu(tr("&Help"))

        add_guide_action(self, h, "pdf-basics")
        self.act(tr("&About"), lambda: QMessageBox.about(
            self, APP_NAME, tr("<b>{APP_NAME}</b><br>Part of {SUITE_NAME}.").format(APP_NAME=APP_NAME, SUITE_NAME=SUITE_NAME)), None, h)

        tb = QToolBar(tr("Main"))
        tb.setObjectName("main")
        tb.setToolButtonStyle(Qt.ToolButtonTextOnly)
        self.addToolBar(tb)
        for a in (a_new, a_open, a_save, a_print):
            tb.addAction(a)
        tb.addSeparator()
        tb.addAction(a_undo)
        tb.addAction(a_redo)
        tb.addSeparator()
        tb.addAction(a_zout)
        self.zoom_box = QComboBox()
        self.zoom_box.setEditable(True)
        self.zoom_box.addItems(["50%", "75%", "100%", "125%", "150%", "200%", "300%"])
        self.zoom_box.setFixedWidth(80)
        self.zoom_box.lineEdit().returnPressed.connect(self._zoom_from_box)
        self.zoom_box.activated.connect(self._zoom_from_box)
        tb.addWidget(self.zoom_box)
        tb.addAction(a_zin)
        tb.addAction(a_fit)
        tb.addSeparator()
        tb.addWidget(QLabel(tr(" Page ")))
        self.page_spin = QSpinBox()
        self.page_spin.setMinimum(1)
        self.page_spin.editingFinished.connect(lambda: self.goto_page(self.page_spin.value() - 1))
        tb.addWidget(self.page_spin)
        self.page_count_label = QLabel(" / 1 ")
        tb.addWidget(self.page_count_label)
        tb.addSeparator()
        self.search_box = QLineEdit()
        self.search_box.setPlaceholderText(tr("Find text..."))
        self.search_box.setFixedWidth(200)
        self.search_box.returnPressed.connect(lambda: self.find(1))
        tb.addWidget(self.search_box)
        self.act(tr("Previous"), lambda: self.find(-1), "Shift+F3", None)
        tb.addAction(self.act(tr("Prev"), lambda: self.find(-1)))
        tb.addAction(self.act(tr("Next"), lambda: self.find(1), "F3"))

        self.addToolBarBreak()
        tt = QToolBar(tr("Tools"))
        tt.setObjectName("tools")
        tt.setToolButtonStyle(Qt.ToolButtonTextOnly)
        self.addToolBar(tt)
        group = QActionGroup(self)
        self.tool_actions = {}
        for key, label, tip in TOOLS:
            a = QAction(label, self, checkable=True)
            a.setToolTip(tip)
            a.setStatusTip(tip)
            a.triggered.connect(lambda _=False, k=key: self.set_tool(k))
            group.addAction(a)
            tt.addAction(a)
            self.tool_actions[key] = a
        self.tool_actions["select"].setChecked(True)
        tt.addSeparator()
        tt.addWidget(QLabel(tr(" Colour ")))
        self.color_btn = ColorButton(QColor(220, 30, 30), tr("Pen / text colour"))
        tt.addWidget(self.color_btn)
        tt.addWidget(QLabel(tr(" Highlight ")))
        self.hl_btn = ColorButton(QColor(255, 230, 0), tr("Highlight colour"))
        tt.addWidget(self.hl_btn)
        tt.addWidget(QLabel(tr(" Width ")))
        self.width_spin = QSpinBox()
        self.width_spin.setRange(1, 20)
        self.width_spin.setValue(2)
        tt.addWidget(self.width_spin)
        tt.addWidget(QLabel(tr(" Font size ")))
        self.font_spin = QSpinBox()
        self.font_spin.setRange(4, 144)
        self.font_spin.setValue(12)
        tt.addWidget(self.font_spin)

    def set_tool(self, key):
        self.tool = key
        self.tool_actions[key].setChecked(True)
        self.sel_words = None
        cursor = Qt.ArrowCursor if key == "select" else (
            Qt.IBeamCursor if key in TEXT_TOOLS or key in ("edittext", "addtext") else Qt.CrossCursor)
        for w in self.page_widgets:
            w.setCursor(cursor)
            w.update()
        if key == "signature" and not self.signature:
            self.create_signature()

    # ------------------------------------------------------------ document lifecycle
    def maybe_save(self):
        if not self.dirty:
            return True
        r = QMessageBox.question(self, APP_NAME, tr("Save changes to this document?"),
                                 QMessageBox.Save | QMessageBox.Discard | QMessageBox.Cancel)
        if r == QMessageBox.Save:
            return self.save()
        return r == QMessageBox.Discard

    def set_document(self, doc, path):
        if self.doc is not None:
            self.doc.close()
        self.doc, self.path = doc, path
        self.dirty = False
        self.undo_stack.clear()
        self.redo_stack.clear()
        self.search_hits, self.search_query = [], ""
        self.rebuild_pages()
        self.scroll.verticalScrollBar().setValue(0)

    def new_blank(self, confirm=True):
        if confirm and not self.maybe_save():
            return
        doc = pymupdf.open()
        doc.new_page(width=612, height=792)
        self.password = None
        self.set_document(doc, None)

    def open_dialog(self):
        f, _ = QFileDialog.getOpenFileName(self, tr("Open PDF"), "", tr("PDF files (*.pdf);;All files (*)"))
        if f:
            self.open_path(f)

    def open_path(self, path):
        if not self.maybe_save():
            return
        try:
            with open(path, "rb") as fh:
                data = fh.read()
            doc = pymupdf.open(stream=data, filetype="pdf")
        except Exception as ex:
            QMessageBox.critical(self, APP_NAME, tr("Could not open file:\n{ex}").format(ex=ex))
            return
        self.password = None
        if doc.needs_pass:
            while True:
                pw, ok = QInputDialog.getText(self, APP_NAME, tr("This PDF is password protected.\nPassword:"), QLineEdit.Password)
                if not ok:
                    return
                if doc.authenticate(pw):
                    self.password = pw
                    break
        self.set_document(doc, path)

    def load_bytes(self, data):
        doc = pymupdf.open(stream=data, filetype="pdf")
        if doc.needs_pass and self.password:
            doc.authenticate(self.password)
        self.doc.close()
        self.doc = doc
        self.rebuild_pages()

    def save(self):
        if not self.path:
            return self.save_as()
        return self.write(self.path)

    def save_as(self):
        start = self.path or "Untitled.pdf"
        f, _ = QFileDialog.getSaveFileName(self, tr("Save PDF As"), start, tr("PDF files (*.pdf)"))
        if not f:
            return False
        if not f.lower().endswith(".pdf"):
            f += ".pdf"
        return self.write(f)

    def write(self, path):
        try:
            data = self.doc.tobytes(garbage=3, deflate=True)
            if self.font_cache:  # we embedded whole system fonts; keep only the glyphs used
                try:
                    tmp = pymupdf.open(stream=data, filetype="pdf")
                    tmp.subset_fonts()
                    data = tmp.tobytes(garbage=3, deflate=True)
                    tmp.close()
                except Exception:
                    pass
            with open(path, "wb") as fh:
                fh.write(data)
        except Exception as ex:
            QMessageBox.critical(self, APP_NAME, tr("Could not save:\n{ex}").format(ex=ex))
            return False
        self.path = path
        self.dirty = False
        self.update_title()
        self.status.showMessage(tr("Saved {path}").format(path=path), 4000)
        return True

    def closeEvent(self, e):
        if self.maybe_save():
            e.accept()
        else:
            e.ignore()

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            e.acceptProposedAction()

    def dropEvent(self, e):
        for url in e.mimeData().urls():
            p = url.toLocalFile()
            if p.lower().endswith(".pdf"):
                self.open_path(p)
                return

    def update_title(self):
        name = os.path.basename(self.path) if self.path else tr("Untitled")
        self.setWindowTitle(f"{'*' if self.dirty else ''}{name} - {APP_NAME}")

    # ------------------------------------------------------------ change tracking / undo
    def begin_change(self):
        self.undo_stack.append(self.doc.tobytes())
        del self.undo_stack[:-40]
        self.redo_stack.clear()

    def end_change(self, idxs=(), structure=False):
        self.dirty = True
        if structure:
            self.rebuild_pages()
        else:
            for i in idxs:
                self.img_cache.pop(i, None)
                self.field_cache.pop(i, None)
                self.page_widgets[i].update()
                self.update_thumb(i)
        self.update_title()

    def undo(self):
        if not self.undo_stack:
            return
        self.redo_stack.append(self.doc.tobytes())
        self.load_bytes(self.undo_stack.pop())
        self.dirty = True
        self.update_title()

    def redo(self):
        if not self.redo_stack:
            return
        self.undo_stack.append(self.doc.tobytes())
        self.load_bytes(self.redo_stack.pop())
        self.dirty = True
        self.update_title()

    # ------------------------------------------------------------ rendering
    def rebuild_pages(self):
        cur = self.current_page() if self.page_widgets else 0
        self.img_cache.clear()
        self.field_cache.clear()
        self.sel_annot = None
        self.sel_words = None
        self.drag = None
        for w in self.page_widgets:
            self.vbox.removeWidget(w)
            w.deleteLater()
        self.page_widgets = []
        for i in range(self.doc.page_count):
            w = PageWidget(self, i)
            self.vbox.addWidget(w, alignment=Qt.AlignHCenter)
            self.page_widgets.append(w)
        self.set_tool(self.tool if self.tool != "signature" or self.signature else "select")
        n = self.doc.page_count
        self.page_spin.setMaximum(n)
        self.page_count_label.setText(f" / {n} ")
        self.zoom_box.setEditText(f"{round(self.zoom * 100)}%")
        self.thumbs.clear()
        for i in range(n):
            it = QListWidgetItem(str(i + 1))
            it.setTextAlignment(Qt.AlignHCenter)
            self.thumbs.addItem(it)
        self.thumb_queue = list(range(n))
        self.thumb_timer.start(0)
        self.update_title()
        QTimer.singleShot(0, lambda: self.goto_page(min(cur, n - 1)))

    def page_image(self, idx):
        img = self.img_cache.get(idx)
        if img is None:
            dpr = self.devicePixelRatioF()
            pix = self.doc[idx].get_pixmap(matrix=pymupdf.Matrix(self.zoom * dpr, self.zoom * dpr),
                                           annots=True)
            img = pix_to_qimage(pix)
            img.setDevicePixelRatio(dpr)
            self.img_cache[idx] = img
        return img

    def _render_some_thumbs(self):
        for _ in range(8):
            if not self.thumb_queue:
                self.thumb_timer.stop()
                return
            self.update_thumb(self.thumb_queue.pop(0))

    def update_thumb(self, i):
        if i >= self.thumbs.count():
            return
        page = self.doc[i]
        s = 130 / max(page.rect.width, page.rect.height) * 1.3
        pix = page.get_pixmap(matrix=pymupdf.Matrix(s, s), annots=True)
        self.thumbs.item(i).setIcon(QIcon(QPixmap.fromImage(pix_to_qimage(pix))))

    def set_zoom(self, z):
        z = max(0.1, min(8.0, z))
        cur = self.current_page()
        self.zoom = z
        self.img_cache.clear()
        for w in self.page_widgets:
            w.resize_to_page()
        self.zoom_box.setEditText(f"{round(z * 100)}%")
        QTimer.singleShot(0, lambda: self.goto_page(cur))

    def _zoom_from_box(self, *_):
        try:
            self.set_zoom(float(self.zoom_box.currentText().strip().rstrip("%")) / 100)
        except ValueError:
            pass

    def fit_width(self):
        if not self.doc.page_count:
            return
        w = max(p.rect.width for p in self.doc)
        self.set_zoom((self.scroll.viewport().width() - 60) / w)

    def toggle_fields(self, on):
        self.show_fields = on
        for w in self.page_widgets:
            w.update()

    def current_page(self):
        if not self.page_widgets:
            return 0
        y = self.scroll.verticalScrollBar().value() + self.scroll.viewport().height() // 3
        for i, w in enumerate(self.page_widgets):
            if w.y() + w.height() + 7 >= y:
                return i
        return len(self.page_widgets) - 1

    def _track_current_page(self):
        self.page_spin.blockSignals(True)
        self.page_spin.setValue(self.current_page() + 1)
        self.page_spin.blockSignals(False)

    def goto_page(self, i):
        if 0 <= i < len(self.page_widgets):
            self.scroll.verticalScrollBar().setValue(self.page_widgets[i].y() - 10)
            self._track_current_page()

    # ------------------------------------------------------------ coordinates
    def to_pdf(self, idx, qp):
        return pymupdf.Point(qp.x() / self.zoom, qp.y() / self.zoom) * self.doc[idx].derotation_matrix

    def to_screen(self, idx, rect):
        r = (pymupdf.Rect(rect) * self.doc[idx].rotation_matrix).normalize()
        z = self.zoom
        return QRectF(r.x0 * z, r.y0 * z, r.width * z, r.height * z)

    # ------------------------------------------------------------ overlay painting
    def field_rects(self, idx):
        if idx not in self.field_cache:
            self.field_cache[idx] = [w.rect for w in self.doc[idx].widgets()]
        return self.field_cache[idx]

    def paint_overlay(self, idx, p):
        if self.show_fields:
            for r in self.field_rects(idx):
                p.fillRect(self.to_screen(idx, r), QColor(60, 130, 255, 40))
        for k, (i, r) in enumerate(self.search_hits):
            if i == idx:
                c = QColor(255, 140, 0, 130) if k == self.search_pos else QColor(255, 220, 0, 90)
                p.fillRect(self.to_screen(idx, r), c)
        if self.sel_words and self.sel_words[0] == idx:
            c = QColor(self.hl_btn.color()) if self.tool == "highlight" else QColor(46, 117, 214)
            c.setAlpha(90)
            for r in line_rects(self.sel_words[1]):
                p.fillRect(self.to_screen(idx, r), c)
        d = self.drag
        if self.sel_annot and self.sel_annot[0] == idx:
            r = self.to_screen(idx, self.sel_annot[2]).adjusted(-3, -3, 3, 3)
            if d and d["mode"] == "move":
                r.translate(d["cur"] - d["start"])
            p.setPen(QPen(QColor(46, 117, 214), 1.5, Qt.DashLine))
            p.setBrush(Qt.NoBrush)
            p.drawRect(r)
        if not d or d["idx"] != idx:
            return
        mode, a, b = d["mode"], d["start"], d["cur"]
        pen = QPen(self.color_btn.color(), max(1.0, self.width_spin.value() * self.zoom))
        p.setPen(pen)
        p.setBrush(Qt.NoBrush)
        rect = QRectF(a, b).normalized()
        if mode == "rect":
            p.drawRect(rect)
        elif mode == "ellipse":
            p.drawEllipse(rect)
        elif mode in ("line", "arrow"):
            p.drawLine(a, b)
        elif mode == "whiteout":
            p.fillRect(rect, Qt.white)
            p.setPen(QPen(QColor("#999"), 1, Qt.DashLine))
            p.drawRect(rect)
        elif mode == "redact":
            p.fillRect(rect, QColor(0, 0, 0, 160))
        elif mode == "image":
            p.setPen(QPen(QColor("#2e75d6"), 1, Qt.DashLine))
            p.drawRect(rect)
        elif mode == "ink" and len(d["pts"]) > 1:
            path = QPainterPath(d["pts"][0])
            for pt in d["pts"][1:]:
                path.lineTo(pt)
            pen.setCapStyle(Qt.RoundCap)
            pen.setJoinStyle(Qt.RoundJoin)
            p.setPen(pen)
            p.drawPath(path)

    # ------------------------------------------------------------ mouse handling
    # Annotations/widgets are only valid while their Page object lives, so keep it referenced.
    def annot_at(self, idx, pt):
        self._held_page = page = self.doc[idx]
        for a in reversed(list(page.annots())):
            if pt in a.rect + (-3, -3, 3, 3):
                return a
        return None

    def widget_at(self, idx, pt):
        self._held_page = page = self.doc[idx]
        for w in page.widgets():
            if pt in w.rect:
                return w
        return None

    def select_annot(self, idx, annot=None):
        old = self.sel_annot
        self.sel_annot = (idx, annot.xref, pymupdf.Rect(annot.rect)) if annot else None
        for i in {old[0] if old else None, idx}:
            if i is not None and i < len(self.page_widgets):
                self.page_widgets[i].update()

    def get_sel_annot(self):
        if not self.sel_annot:
            return None, None
        page = self.doc[self.sel_annot[0]]
        return page, page.load_annot(self.sel_annot[1])

    def on_press(self, pw, e):
        if e.button() != Qt.LeftButton:
            return
        idx, pos = pw.idx, e.position()
        self.drag = {"idx": idx, "start": pos, "cur": pos, "pts": [pos], "mode": self.tool}
        pt = self.to_pdf(idx, pos)
        if self.tool == "select":
            a = self.annot_at(idx, pt)
            if a:
                self.select_annot(idx, a)
                self.drag["mode"] = "move"
                return
            if self.sel_annot:
                self.select_annot(None)
            w = self.widget_at(idx, pt)
            if w:
                self.drag = None
                self.fill_field(idx, w)
                return
            self.drag["mode"] = "textsel"
        if self.drag["mode"] in TEXT_TOOLS + ("textsel",):
            self.drag["words"] = self.doc[idx].get_text("words")
            if self.sel_words:
                old = self.sel_words[0]
                self.sel_words = None
                self.page_widgets[old].update()

    def on_move(self, pw, e):
        d = self.drag
        if not d:
            return
        d["cur"] = e.position()
        d["pts"].append(e.position())
        if d["mode"] in TEXT_TOOLS + ("textsel",):
            sel = words_between(d["words"], self.to_pdf(d["idx"], d["start"]),
                                self.to_pdf(d["idx"], d["cur"]))
            self.sel_words = (d["idx"], sel) if sel else None
        pw.update()

    def on_release(self, pw, e):
        d = self.drag
        if not d or e.button() != Qt.LeftButton:
            return
        self.drag = None
        idx, mode = d["idx"], d["mode"]
        delta = d["cur"] - d["start"]
        moved = abs(delta.x()) + abs(delta.y()) > 4
        p1, p2 = self.to_pdf(idx, d["start"]), self.to_pdf(idx, d["cur"])
        rect = pymupdf.Rect(p1, p2).normalize()
        try:
            if mode == "move":
                if moved:
                    self.move_annot(p2 - p1)
            elif mode == "textsel":
                if not moved:
                    self.sel_words = None
                elif self.sel_words:
                    self.status.showMessage(tr("Text selected - press Ctrl+C to copy"), 3000)
            elif mode in TEXT_TOOLS:
                if moved and self.sel_words:
                    self.add_markup(idx, mode, self.sel_words[1])
                self.sel_words = None
            elif mode in DRAG_TOOLS:
                if moved:
                    self.add_shape(idx, mode, p1, p2, rect)
            elif mode == "ink":
                if len(d["pts"]) > 1:
                    self.add_ink(idx, [self.to_pdf(idx, q) for q in d["pts"]])
            elif mode == "addtext":
                self.add_text(idx, p2)
            elif mode == "edittext":
                self.edit_text_at(idx, p2)
            elif mode == "note":
                self.add_note(idx, p2)
            elif mode == "image":
                self.add_image(idx, rect if moved else None, p2)
            elif mode == "signature":
                self.place_signature(idx, d["cur"])
        except Exception as ex:
            QMessageBox.warning(self, APP_NAME, tr("That didn't work:\n{ex}").format(ex=ex))
        pw.update()

    def on_double_click(self, pw, e):
        if self.tool != "select":
            return
        a = self.annot_at(pw.idx, self.to_pdf(pw.idx, e.position()))
        if a:
            self.select_annot(pw.idx, a)
            self.edit_annot_text()

    def on_context_menu(self, pw, e):
        pt = self.to_pdf(pw.idx, e.pos())
        a = self.annot_at(pw.idx, pt)
        m = QMenu(self)
        if a:
            self.select_annot(pw.idx, a)
            if a.type[0] in (pymupdf.PDF_ANNOT_FREE_TEXT, pymupdf.PDF_ANNOT_TEXT):
                m.addAction(tr("Edit Text..."), self.edit_annot_text)
            m.addAction(tr("Change Colour..."), self.recolor_annot)
            m.addAction("Delete", self.delete_selected_annot)
        else:
            if self.sel_words:
                m.addAction(tr("Copy"), self.copy_text)
            m.addMenu(self.pages_menu)
        m.exec(e.globalPos())

    # ------------------------------------------------------------ editing operations
    def add_markup(self, idx, kind, words):
        page = self.doc[idx]
        self.begin_change()
        f = {"highlight": page.add_highlight_annot, "underline": page.add_underline_annot,
             "strike": page.add_strikeout_annot}[kind]
        a = f(line_rects(words))
        a.set_colors(stroke=rgb(self.hl_btn.color() if kind == "highlight" else self.color_btn.color()))
        a.update()
        self.end_change([idx])

    def add_shape(self, idx, kind, p1, p2, rect):
        page = self.doc[idx]
        self.begin_change()
        if kind == "whiteout":
            page.draw_rect(rect, color=None, fill=(1, 1, 1), overlay=True)
        elif kind == "redact":
            page.add_redact_annot(rect, fill=(0, 0, 0))
            page.apply_redactions()
        else:
            c = rgb(self.color_btn.color())
            if kind == "rect":
                a = page.add_rect_annot(rect)
            elif kind == "ellipse":
                a = page.add_circle_annot(rect)
            else:
                a = page.add_line_annot(p1, p2)
                if kind == "arrow":
                    a.set_line_ends(pymupdf.PDF_ANNOT_LE_NONE, pymupdf.PDF_ANNOT_LE_CLOSED_ARROW)
                    a.set_colors(stroke=c, fill=c)
            if kind != "arrow":
                a.set_colors(stroke=c)
            a.set_border(width=self.width_spin.value())
            a.update()
        self.end_change([idx])

    def add_ink(self, idx, pts, color=None, width=None):
        page = self.doc[idx]
        self.begin_change()
        strokes = pts if pts and isinstance(pts[0], list) else [pts]
        a = page.add_ink_annot([[(p.x, p.y) for p in s] for s in strokes])
        a.set_colors(stroke=rgb(color or self.color_btn.color()))
        a.set_border(width=width or self.width_spin.value())
        a.update()
        self.end_change([idx])

    def _freetext(self, page, origin, text, size, family, color):
        code, _, ft = FONTS[family]
        lines = text.split("\n")
        w = max(pymupdf.get_text_length(line, fontname=code, fontsize=size) for line in lines)
        rect = pymupdf.Rect(origin.x, origin.y, origin.x + w + size,
                            origin.y + len(lines) * size * 1.25 + size * 0.6)
        page.add_freetext_annot(rect, text, fontsize=size, fontname=ft, text_color=rgb(color))

    def load_font(self, family, bold, italic):
        """Installed font as a pymupdf.Font (cached); falls back to the built-in Helvetica."""
        key = (family, bold, italic)
        if key not in self.font_cache:
            path = font_file(family, bold, italic)
            try:
                self.font_cache[key] = pymupdf.Font(fontfile=path) if path else None
            except Exception:
                self.font_cache[key] = None
            if self.font_cache[key] is None:
                self.font_cache[key] = pymupdf.Font(
                    {(0, 0): "helv", (1, 0): "hebo", (0, 1): "heit", (1, 1): "hebi"}[(bold, italic)])
        return self.font_cache[key]

    def write_text(self, page, baseline, text, size, family, bold, italic, color):
        """Write real page text (not an annotation) using an installed font."""
        font = self.load_font(family, bold, italic)
        tw = pymupdf.TextWriter(page.rect)
        x, y = baseline
        for line in text.split("\n"):
            tw.append((x, y), line, font=font, fontsize=size)
            y += size * 1.2
        tw.write_text(page, color=rgb(color))

    def add_text(self, idx, pt):
        dlg = TextDialog(self, tr("Add Text"), size=self.font_spin.value(), color=self.color_btn.color())
        if not dlg.exec():
            return
        text, size, family, color, bold, italic = dlg.values()
        if not text.strip():
            return
        self.begin_change()
        asc = self.load_font(family, bold, italic).ascender or 0.8
        self.write_text(self.doc[idx], (pt.x, pt.y + size * asc), text, size, family, bold, italic, color)
        self.end_change([idx])

    def edit_annot_text(self):
        page, a = self.get_sel_annot()
        if not a:
            return
        if a.type[0] == pymupdf.PDF_ANNOT_TEXT:
            text, ok = QInputDialog.getMultiLineText(self, tr("Edit Note"), tr("Note:"), a.info["content"])
            if ok:
                self.begin_change()
                a.set_info(content=text)
                a.update()
                self.end_change([self.sel_annot[0]])
            return
        if a.type[0] != pymupdf.PDF_ANNOT_FREE_TEXT:
            return
        da = self.doc.xref_get_key(a.xref, "DA")[1]
        size, family, color = 12, "Helvetica", QColor("black")
        m = re.search(r"/(\w+)\s+([\d.]+)\s+Tf", da)
        if m:
            family = FREETEXT_FONT.get(m.group(1), "Helvetica")
            size = float(m.group(2))
        m = re.search(r"([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+rg", da)
        if m:
            color = QColor.fromRgbF(*(float(x) for x in m.groups()))
        dlg = TextDialog(self, tr("Edit Text Box"), a.info["content"], size, family, color, styles=False)
        if not dlg.exec():
            return
        text, size, family, color, _, _ = dlg.values()
        idx, origin = self.sel_annot[0], a.rect.tl
        self.begin_change()
        page.delete_annot(a)
        if text.strip():
            self._freetext(page, origin, text, size, family, color)
        self.sel_annot = None
        self.end_change([idx])

    def recolor_annot(self):
        page, a = self.get_sel_annot()
        if not a:
            return
        from PySide6.QtWidgets import QColorDialog
        c = QColorDialog.getColor(self.color_btn.color(), self, tr("Annotation colour"))
        if not c.isValid():
            return
        idx = self.sel_annot[0]
        self.begin_change()
        if a.type[0] == pymupdf.PDF_ANNOT_FREE_TEXT:
            a.update(text_color=rgb(c))
        else:
            fill = rgb(c) if a.colors.get("fill") else None
            a.set_colors(stroke=rgb(c), fill=fill)
            a.update()
        self.end_change([idx])

    def delete_selected_annot(self):
        page, a = self.get_sel_annot()
        if not a:
            return
        idx = self.sel_annot[0]
        self.begin_change()
        page.delete_annot(a)
        self.sel_annot = None
        self.end_change([idx])

    def move_annot(self, delta):
        page, a = self.get_sel_annot()
        if not a:
            return
        idx = self.sel_annot[0]
        dx, dy = delta.x, delta.y
        self.begin_change()
        t = a.type[0]
        if t in (pymupdf.PDF_ANNOT_INK, pymupdf.PDF_ANNOT_LINE):
            verts, colors, width = a.vertices, a.colors, a.border.get("width", 1) or 1
            ends = a.line_ends if t == pymupdf.PDF_ANNOT_LINE else None
            page.delete_annot(a)
            if t == pymupdf.PDF_ANNOT_INK:
                new = page.add_ink_annot([[(x + dx, y + dy) for x, y in s] for s in verts])
            else:
                (x0, y0), (x1, y1) = verts[0], verts[-1]
                new = page.add_line_annot((x0 + dx, y0 + dy), (x1 + dx, y1 + dy))
                if ends:
                    new.set_line_ends(*ends)
            new.set_colors(stroke=colors.get("stroke") or None, fill=colors.get("fill") or None)
            new.set_border(width=width)
            new.update()
        else:
            a.set_rect(a.rect + (dx, dy, dx, dy))
            new = a
        self.end_change([idx])
        self.select_annot(idx, page.load_annot(new.xref))

    def add_note(self, idx, pt):
        text, ok = QInputDialog.getMultiLineText(self, tr("Sticky Note"), tr("Note:"))
        if ok and text.strip():
            self.begin_change()
            a = self.doc[idx].add_text_annot(pt, text)
            a.set_colors(stroke=rgb(self.hl_btn.color()))
            a.update()
            self.end_change([idx])

    def add_image(self, idx, rect, pt):
        f, _ = QFileDialog.getOpenFileName(self, tr("Insert Image"), "",
                                           tr("Images (*.png *.jpg *.jpeg *.bmp *.gif *.tif *.tiff)"))
        if not f:
            return
        if rect is None:
            img = QImage(f)
            if img.isNull():
                raise ValueError(tr("Unsupported image"))
            w = 200
            rect = pymupdf.Rect(pt.x, pt.y, pt.x + w, pt.y + w * img.height() / img.width())
        self.begin_change()
        self.doc[idx].insert_image(rect, filename=f, keep_proportion=True)
        self.end_change([idx])

    def edit_text_at(self, idx, pt):
        page = self.doc[idx]
        found = None
        for b in page.get_text("dict")["blocks"]:
            if b["type"] != 0:
                continue
            for line in b["lines"]:
                if line["spans"] and pt in pymupdf.Rect(line["bbox"]) + (-2, -2, 2, 2):
                    found = line
                    break
            if found:
                break
        if not found:
            self.status.showMessage(tr("No text there - click directly on a line of text"), 4000)
            return
        spans = [s for s in found["spans"] if s["text"].strip()] or found["spans"]
        s0 = spans[0]
        text = "".join(s["text"] for s in found["spans"]).strip()
        name = s0["font"].lower()
        family = match_family(s0["font"])
        if not family:  # font not installed: pick a look-alike
            if any(k in name for k in ("times", "serif", "roman", "georgia", "garamond", "book")):
                family = "Times New Roman"
            elif any(k in name for k in ("cour", "mono", "consol")):
                family = "Courier New"
            else:
                family = "Arial"
        bold = bool(s0["flags"] & 16) or "bold" in name
        italic = bool(s0["flags"] & 2) or "italic" in name or "oblique" in name
        dlg = TextDialog(self, tr("Edit Text"), text, s0["size"], family, int_to_qcolor(s0["color"]),
                         bold, italic)
        if not dlg.exec():
            return
        new_text, size, family, color, bold, italic = dlg.values()
        self.begin_change()
        r = pymupdf.Rect(found["bbox"])
        h = r.height
        r.y0 += h * 0.2
        r.y1 -= h * 0.2
        page.add_redact_annot(r, fill=False)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                              graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        if new_text.strip():
            x0 = min(s["origin"][0] for s in spans)
            self.write_text(page, (x0, s0["origin"][1]), new_text, size, family, bold, italic, color)
        self.end_change([idx])

    def create_signature(self):
        dlg = SignatureDialog(self)
        if dlg.exec() and dlg.result_sig:
            self.signature = dlg.result_sig
            self.set_tool("signature")
            self.status.showMessage(tr("Signature ready - click on a page to place it"), 5000)
            return True
        if self.tool == "signature" and not self.signature:
            self.set_tool("select")
        return False

    def place_signature(self, idx, screen_pt):
        if not self.signature:
            return
        kind, data = self.signature
        width = 170 * self.zoom
        if kind == "ink":
            xs = [x for s in data for x, _ in s]
            ys = [y for s in data for _, y in s]
            bx, by = min(xs), min(ys)
            bw, bh = max(max(xs) - bx, 1), max(max(ys) - by, 1)
            sc = width / bw
            ox, oy = screen_pt.x() - width / 2, screen_pt.y() - bh * sc / 2
            strokes = [[self.to_pdf(idx, QPointF(ox + (x - bx) * sc, oy + (y - by) * sc))
                        for x, y in s] for s in data]
            self.add_ink(idx, strokes, QColor("#102a6b"), 2)
        else:
            img = QImage(data)
            if img.isNull():
                raise ValueError(tr("Could not read signature image"))
            h = width * img.height() / img.width()
            r = pymupdf.Rect(self.to_pdf(idx, screen_pt - QPointF(width / 2, h / 2)),
                             self.to_pdf(idx, screen_pt + QPointF(width / 2, h / 2))).normalize()
            self.begin_change()
            self.doc[idx].insert_image(r, filename=data, keep_proportion=True)
            self.end_change([idx])

    def fill_field(self, idx, w):
        ft = w.field_type
        label = w.field_label or w.field_name or tr("Field")
        if ft in (pymupdf.PDF_WIDGET_TYPE_CHECKBOX, pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON):
            checked = w.field_value not in (False, "Off", "", None)
            value = True if ft == pymupdf.PDF_WIDGET_TYPE_RADIOBUTTON else not checked
        elif ft in (pymupdf.PDF_WIDGET_TYPE_COMBOBOX, pymupdf.PDF_WIDGET_TYPE_LISTBOX):
            choices = [c[1] if isinstance(c, (list, tuple)) else c for c in (w.choice_values or [])]
            if not choices:
                return
            cur = choices.index(w.field_value) if w.field_value in choices else 0
            value, ok = QInputDialog.getItem(self, tr("Fill Field"), label, choices, cur, False)
            if not ok:
                return
        elif ft == pymupdf.PDF_WIDGET_TYPE_TEXT:
            cur = w.field_value or ""
            if w.field_flags & pymupdf.PDF_TX_FIELD_IS_MULTILINE:
                value, ok = QInputDialog.getMultiLineText(self, tr("Fill Field"), label, cur)
            else:
                value, ok = QInputDialog.getText(self, tr("Fill Field"), label, QLineEdit.Normal, cur)
            if not ok:
                return
        else:
            return
        self.begin_change()
        w.field_value = value
        w.update()
        self.end_change([idx])

    def copy_text(self):
        if self.sel_words:
            QGuiApplication.clipboard().setText(words_text(self.sel_words[1]))
            self.status.showMessage(tr("Copied"), 2000)

    # ------------------------------------------------------------ search
    def find(self, direction):
        q = self.search_box.text().strip()
        if not q:
            self.search_hits = []
            for w in self.page_widgets:
                w.update()
            return
        if q != self.search_query:
            self.search_query = q
            self.search_hits = [(i, r) for i, page in enumerate(self.doc) for r in page.search_for(q)]
            start = self.current_page()
            self.search_pos = next((k - direction for k, (i, _) in enumerate(self.search_hits)
                                    if i >= start), -1)
            for w in self.page_widgets:
                w.update()
        if not self.search_hits:
            self.status.showMessage(tr('"{q}" not found').format(q=q), 3000)
            return
        old = self.search_hits[self.search_pos][0] if 0 <= self.search_pos < len(self.search_hits) else None
        self.search_pos = (self.search_pos + direction) % len(self.search_hits)
        idx, r = self.search_hits[self.search_pos]
        for i in {old, idx} - {None}:
            self.page_widgets[i].update()
        sr = self.to_screen(idx, r)
        self.scroll.ensureVisible(int(sr.center().x() + self.page_widgets[idx].x()),
                                  int(sr.center().y() + self.page_widgets[idx].y()), 50, 150)
        self.status.showMessage(tr("Match {v} of {v2}").format(v=self.search_pos + 1, v2=len(self.search_hits)), 3000)

    # ------------------------------------------------------------ page operations
    def selected_pages(self):
        rows = sorted({self.thumbs.row(it) for it in self.thumbs.selectedItems()})
        return rows or [self.current_page()]

    def rotate_pages(self, deg):
        pages = self.selected_pages()
        self.begin_change()
        for i in pages:
            page = self.doc[i]
            page.set_rotation((page.rotation + deg) % 360)
        self.end_change(structure=True)

    def delete_pages(self):
        pages = self.selected_pages()
        if len(pages) >= self.doc.page_count:
            QMessageBox.information(self, APP_NAME, tr("A PDF needs at least one page."))
            return
        if QMessageBox.question(self, APP_NAME, tr("Delete {v} page(s)?").format(v=len(pages))) != QMessageBox.Yes:
            return
        self.begin_change()
        self.doc.delete_pages(pages)
        self.end_change(structure=True)

    def move_page(self, step):
        i = self.selected_pages()[0]
        j = i + step
        if not 0 <= j < self.doc.page_count:
            return
        self.begin_change()
        self.doc.move_page(i, j if step < 0 else (j + 1 if j + 1 < self.doc.page_count else -1))
        self.end_change(structure=True)
        self.thumbs.setCurrentRow(j)
        self.goto_page(j)

    def insert_blank(self):
        i = self.current_page()
        r = self.doc[i].rect
        self.begin_change()
        self.doc.new_page(pno=i + 1, width=r.width, height=r.height)
        self.end_change(structure=True)
        QTimer.singleShot(0, lambda: self.goto_page(i + 1))

    def insert_from_file(self):
        files, _ = QFileDialog.getOpenFileNames(self, tr("Insert Pages From"), "", tr("PDF files (*.pdf)"))
        if not files:
            return
        at = self.current_page() + 1
        self.begin_change()
        for f in files:
            src = pymupdf.open(f)
            n = src.page_count
            self.doc.insert_pdf(src, start_at=at)
            at += n
            src.close()
        self.end_change(structure=True)

    def _images_to_pdf(self, files, target, at=-1):
        for f in files:
            img = pymupdf.open(f)
            pdf = pymupdf.open("pdf", img.convert_to_pdf())
            target.insert_pdf(pdf, start_at=at)
            if at >= 0:
                at += pdf.page_count
            img.close()
            pdf.close()

    def insert_image_pages(self):
        files, _ = QFileDialog.getOpenFileNames(self, tr("Insert Images as Pages"), "",
                                                tr("Images (*.png *.jpg *.jpeg *.bmp *.gif *.tif *.tiff)"))
        if files:
            self.begin_change()
            self._images_to_pdf(files, self.doc, self.current_page() + 1)
            self.end_change(structure=True)

    def new_from_images(self):
        files, _ = QFileDialog.getOpenFileNames(self, tr("Create PDF from Images"), "",
                                                tr("Images (*.png *.jpg *.jpeg *.bmp *.gif *.tif *.tiff)"))
        if not files or not self.maybe_save():
            return
        doc = pymupdf.open()
        self._images_to_pdf(files, doc)
        self.set_document(doc, None)
        self.dirty = True
        self.update_title()

    def combine_files(self):
        files, _ = QFileDialog.getOpenFileNames(
            self, tr("Choose files to combine (in order)"), "",
            tr("PDFs and images (*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff)"))
        if not files or not self.maybe_save():
            return
        doc = pymupdf.open()
        for f in files:
            if f.lower().endswith(".pdf"):
                src = pymupdf.open(f)
                doc.insert_pdf(src)
                src.close()
            else:
                self._images_to_pdf([f], doc)
        self.set_document(doc, None)
        self.dirty = True
        self.update_title()

    def extract_pages(self):
        pages = self.selected_pages()
        f, _ = QFileDialog.getSaveFileName(self, tr("Save Extracted Pages"), "extracted.pdf",
                                           tr("PDF files (*.pdf)"))
        if not f:
            return
        out = pymupdf.open()
        for i in pages:
            out.insert_pdf(self.doc, from_page=i, to_page=i)
        out.save(f, garbage=3, deflate=True)
        out.close()
        self.status.showMessage(tr("Saved {v} page(s) to {f}").format(v=len(pages), f=f), 5000)

    def split_pages(self):
        folder = QFileDialog.getExistingDirectory(self, tr("Folder for the single pages"))
        if not folder:
            return
        base = os.path.splitext(os.path.basename(self.path or "document.pdf"))[0]
        for i in range(self.doc.page_count):
            out = pymupdf.open()
            out.insert_pdf(self.doc, from_page=i, to_page=i)
            out.save(os.path.join(folder, f"{base}_page{i + 1}.pdf"), garbage=3, deflate=True)
            out.close()
        self.status.showMessage(tr("Saved {page_count} files to {folder}").format(page_count=self.doc.page_count, folder=folder), 5000)

    # ------------------------------------------------------------ export / print
    def export_images(self):
        folder = QFileDialog.getExistingDirectory(self, tr("Folder for the page images"))
        if not folder:
            return
        fmt, ok = QInputDialog.getItem(self, tr("Export Images"), tr("Format:"), ["PNG", "JPG"], 0, False)
        if not ok:
            return
        dpi, ok = QInputDialog.getInt(self, tr("Export Images"), tr("Resolution (DPI):"), 150, 50, 600)
        if not ok:
            return
        base = os.path.splitext(os.path.basename(self.path or "document.pdf"))[0]
        for i, page in enumerate(self.doc):
            page.get_pixmap(dpi=dpi, annots=True).save(
                os.path.join(folder, f"{base}_page{i + 1}.{fmt.lower()}"))
        self.status.showMessage(tr("Exported {page_count} images to {folder}").format(page_count=self.doc.page_count, folder=folder), 5000)

    def export_text(self):
        f, _ = QFileDialog.getSaveFileName(self, tr("Export Text"), "document.txt", tr("Text files (*.txt)"))
        if f:
            with open(f, "w", encoding="utf-8") as fh:
                fh.write("\n\f".join(page.get_text() for page in self.doc))

    def export_docx(self):
        import docx
        from docx.shared import Pt, RGBColor, Inches
        start = os.path.splitext(self.path or "document.pdf")[0] + ".docx"
        f, _ = QFileDialog.getSaveFileName(self, tr("Export to Word"), start, tr("Word documents (*.docx)"))
        if not f:
            return
        QGuiApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            out = docx.Document()
            for i, page in enumerate(self.doc):
                if i:
                    out.add_page_break()
                for b in page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT)["blocks"]:
                    if b["type"] == 1:
                        try:
                            w = min((b["bbox"][2] - b["bbox"][0]) / 72, 6.5)
                            out.add_picture(io.BytesIO(b["image"]), width=Inches(w))
                        except Exception:
                            pass
                        continue
                    para = out.add_paragraph()
                    for li, line in enumerate(b["lines"]):
                        for s in line["spans"]:
                            run = para.add_run(s["text"])
                            run.font.size = Pt(round(s["size"] * 2) / 2)
                            run.bold = bool(s["flags"] & 16)
                            run.italic = bool(s["flags"] & 2)
                            run.font.color.rgb = RGBColor.from_string(f"{s['color']:06X}")
                            run.font.name = s["font"].split("+")[-1].split("-")[0]
                        if li < len(b["lines"]) - 1:
                            para.add_run(" ")
            out.save(f)
        except Exception as ex:
            QMessageBox.critical(self, APP_NAME, tr("Export failed:\n{ex}").format(ex=ex))
            return
        finally:
            QGuiApplication.restoreOverrideCursor()
        self.status.showMessage(tr("Exported to {f}").format(f=f), 5000)

    def print_doc(self):
        printer = QPrinter(QPrinter.HighResolution)
        dlg = QPrintDialog(printer, self)
        dlg.setMinMax(1, self.doc.page_count)
        if not dlg.exec():
            return
        first, last = printer.fromPage(), printer.toPage()
        if first == 0:
            first, last = 1, self.doc.page_count
        painter = QPainter(printer)
        try:
            for n, i in enumerate(range(first - 1, last)):
                if n:
                    printer.newPage()
                img = pix_to_qimage(self.doc[i].get_pixmap(dpi=200, annots=True))
                vp = painter.viewport()
                size = img.size().scaled(vp.size(), Qt.KeepAspectRatio)
                painter.drawImage(QRectF((vp.width() - size.width()) / 2, 0,
                                         size.width(), size.height()), img)
        finally:
            painter.end()


def main():
    app = make_app(APP_NAME, sys.argv)
    win = PdfEditor(sys.argv[1] if len(sys.argv) > 1 else None)
    win.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
